import base64
import hashlib
import html
import json
import re
import time
import unicodedata
from collections import defaultdict
from io import BytesIO
from pathlib import Path

import streamlit as st
from openai import OpenAI

from main import build_claim_ledger
from i18n import (
    choose,
    domain_label,
    flag_type_label,
    language_label,
    review_status_label,
)

PROVIDERS = {
    "DeepSeek": {
        "base_url": "https://api.deepseek.com",
        "default_model": "deepseek-flash",
        "key_label": "DeepSeek API Key",
    },
    "OpenAI": {
        "base_url": None,
        "default_model": "gpt-5.6-luna",
        "key_label": "OpenAI API Key",
    },
    "OpenRouter": {
        "base_url": "https://openrouter.ai/api/v1",
        "default_model": "cohere/north-mini-code:free",
        "key_label": "OpenRouter API Key",
    },
}

EVIDENCE_STATUSES = {
    "Unverified",
    "Supported",
    "Projection",
    "Subjective / Marketing",
}

CATEGORIES = {
    "公司概况",
    "技术与产品",
    "客户与订单",
    "商业化与产能",
    "财务与预测",
    "团队",
    "知识产权",
    "市场与竞争",
    "风险与限制",
    "其他",
}

SUPPORTED_TYPES = ["txt", "pdf", "docx", "pptx"]
BATCH_SIZE = 5
MAX_VISION_UNITS = 2
TRIAGE_MODEL_MIN_UNITS = 41
MATRIX_DOMAIN_ORDER = [
    "公司概览",
    "商业化与竞争",
    "财务与融资",
    "团队与组织",
    "产品与技术",
    "知识产权",
    "其他重要事项",
]
MATRIX_TOPICS_PER_DOMAIN = 2
PIPELINE_VERSION = "2026-09-20-workspace-v5-i18n"
ANALYST_REVIEW_LABELS = {
    "unreviewed": "○ 未复核",
    "cleared": "✓ 已复核 / 暂无异议",
    "follow_up": "⚑ 需进一步跟进",
}
ANALYST_REVIEW_OPTIONS = [
    "unreviewed",
    "cleared",
    "follow_up",
]
RESEARCH_FLAG_TYPES = {
    "unresolved_conflict",
    "unresolved_version",
    "scope_or_definition_gap",
    "material_evidence_gap",
}
RESEARCH_FLAG_TYPE_LABELS = {
    "unresolved_conflict": "信息冲突待确认",
    "unresolved_version": "版本口径待确认",
    "scope_or_definition_gap": "口径 / 定义待确认",
    "material_evidence_gap": "关键证据缺口",
}

st.set_page_config(
    page_title="Investment Research AI",
    page_icon="📄",
    layout="wide",
)


def current_ui_language():
    return st.session_state.get("ui_language", "zh-CN")


def ui(zh_text, en_text):
    return choose(
        current_ui_language(),
        zh_text,
        en_text,
    )


def create_client(provider, api_key):
    config = PROVIDERS[provider]

    if config["base_url"]:
        return OpenAI(
            api_key=api_key,
            base_url=config["base_url"],
        )

    return OpenAI(api_key=api_key)


def reset_run_stats():
    st.session_state["run_stats"] = {
        "api_calls": 0,
        "claim_calls": 0,
        "vision_calls": 0,
        "review_calls": 0,
        "triage_calls": 0,
        "synthesis_calls": 0,
        "verify_calls": 0,
    }


def record_api_call(call_type):
    stats = st.session_state.setdefault(
        "run_stats",
        {
            "api_calls": 0,
            "claim_calls": 0,
            "vision_calls": 0,
            "review_calls": 0,
            "triage_calls": 0,
            "synthesis_calls": 0,
            "verify_calls": 0,
        },
    )
    stats["api_calls"] += 1
    key = f"{call_type}_calls"
    if key in stats:
        stats[key] += 1


def get_run_stats():
    return st.session_state.get(
        "run_stats",
        {
            "api_calls": 0,
            "claim_calls": 0,
            "vision_calls": 0,
            "review_calls": 0,
            "triage_calls": 0,
            "synthesis_calls": 0,
            "verify_calls": 0,
        },
    )


def reset_stage_timings():
    st.session_state["stage_timings"] = {
        "triage": 0.0,
        "extract": 0.0,
        "vision": 0.0,
        "synthesis": 0.0,
        "verify": 0.0,
    }


def add_stage_time(stage, seconds):
    timings = st.session_state.setdefault(
        "stage_timings",
        {
            "triage": 0.0,
            "extract": 0.0,
            "vision": 0.0,
            "synthesis": 0.0,
            "verify": 0.0,
        },
    )
    if stage in timings:
        timings[stage] += max(0.0, float(seconds))


def get_stage_timings():
    return dict(
        st.session_state.get(
            "stage_timings",
            {
                "triage": 0.0,
                "extract": 0.0,
                "vision": 0.0,
                "synthesis": 0.0,
                "verify": 0.0,
            },
        )
    )


def decode_txt(file_bytes):
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return file_bytes.decode("utf-16")


def compact_text(text, limit=5000):
    text = (text or "").strip()
    if len(text) <= limit:
        return text
    return text[:limit] + "\n[上下文过长，已截断显示]"


def make_text_units(text, unit_label="Lines", group_size=8):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    units = []

    for i in range(0, len(lines), group_size):
        current = lines[i:i + group_size]
        start = i + 1
        end = i + len(current)
        before = lines[max(0, i - 2):i]
        after = lines[end:min(len(lines), end + 2)]
        context = "\n".join(before + current + after)
        units.append({
            "source_location": f"{unit_label} {start}-{end}",
            "text": "\n".join(current),
            "source_context": context,
            "visual_evidence": [],
            "image_count": 0,
            "visual_analysis_status": "not_applicable",
        })

    return units


def image_to_data_url(image_bytes, mime_type):
    encoded = base64.b64encode(image_bytes).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def analyze_visual(client, model_name, image_bytes, mime_type, location, nearby_text=""):
    prompt = f"""
你正在分析投资研究材料中的视觉证据，位置是：{location}。

目标不是替投资人下结论，而是忠实读取图片、截图、表格或图表中明确可见的信息。

规则：
1. 只提取图片中明确可读、明确标注的信息。
2. 可以读取标题、标签、数字、表格单元格、图例和注释。
3. 如果柱状图/折线图没有明确数字标注，不要根据高度自行估算精确数值。
4. 不要因为出现客户 Logo 就推断已经签约或产生收入。
5. 不要把模型自己的解释伪装成原文。
6. 如果图片与投资研究无关（装饰图、Logo、小图标），明确返回无有效证据。

附近的页面/幻灯片文字（仅用于理解上下文）：
{compact_text(nearby_text, 1800)}

请严格返回 JSON 对象：
{{
  "summary": "对视觉内容的简短描述",
  "evidence": [
    {{
      "text": "图片中明确可见的文字或数值",
      "kind": "chart|table|screenshot|image_text|other"
    }}
  ]
}}
"""

    record_api_call("vision")
    response = client.responses.create(
        model=model_name,
        input=[
            {
                "role": "user",
                "content": [
                    {"type": "input_text", "text": prompt},
                    {
                        "type": "input_image",
                        "image_url": image_to_data_url(image_bytes, mime_type),
                    },
                ],
            }
        ],
    )

    raw = response.output_text.strip()

    try:
        parsed = json.loads(raw)
        evidence = parsed.get("evidence", [])
        if not isinstance(evidence, list):
            evidence = []
        return {
            "summary": parsed.get("summary", ""),
            "evidence": evidence,
            "status": "analyzed",
        }
    except json.JSONDecodeError:
        return {
            "summary": raw,
            "evidence": [],
            "status": "analyzed_unstructured",
        }


def extract_pdf_units(
    file_bytes,
    analyze_visuals=False,
    client=None,
    model_name=None,
    status_callback=None,
):
    import fitz

    document = fitz.open(stream=file_bytes, filetype="pdf")
    units = []
    total_pages = len(document)

    for page_index, page in enumerate(document, start=1):
        if status_callback:
            status_callback(f"快速解析 PDF：Page {page_index}/{total_pages}")

        text = page.get_text("text").strip()
        image_count = len(page.get_images(full=True))
        drawing_count = len(page.get_drawings())

        preview_pix = page.get_pixmap(
            matrix=fitz.Matrix(0.72, 0.72),
            alpha=False,
        )
        preview_bytes = preview_pix.tobytes(
            "jpeg",
            jpg_quality=58,
        )

        vision_candidate = (
            analyze_visuals
            and (
                len(text) < 180
                or drawing_count >= 12
                or (
                    image_count >= 2
                    and len(text) < 700
                )
            )
        )

        units.append({
            "source_location": f"Page {page_index}",
            "text": text or "[该页未提取到文本]",
            "source_context": text,
            "visual_evidence": [],
            "image_count": image_count,
            "drawing_count": drawing_count,
            "visual_analysis_status": (
                "candidate" if vision_candidate else "not_requested"
            ),
            "vision_candidate": vision_candidate,
            "_vision_images": [(preview_bytes, "image/jpeg")],
            "preview_bytes": preview_bytes,
        })

    return units

def extract_docx_units(file_bytes):
    from docx import Document

    document = Document(BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    units = []
    group_size = 5

    for i in range(0, len(paragraphs), group_size):
        current = paragraphs[i:i + group_size]
        start = i + 1
        end = i + len(current)
        before = paragraphs[max(0, i - 2):i]
        after = paragraphs[end:min(len(paragraphs), end + 2)]
        units.append({
            "source_location": f"Paragraphs {start}-{end}",
            "text": "\n".join(current),
            "source_context": "\n".join(before + current + after),
            "visual_evidence": [],
            "image_count": 0,
            "visual_analysis_status": "not_applicable",
        })

    return units


def extract_pptx_units(
    file_bytes,
    analyze_visuals=False,
    client=None,
    model_name=None,
    status_callback=None,
):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(BytesIO(file_bytes))
    units = []
    total_slides = len(presentation.slides)

    for slide_index, slide in enumerate(presentation.slides, start=1):
        if status_callback:
            status_callback(
                f"快速解析 PPTX：Slide {slide_index}/{total_slides}"
            )

        texts = []
        candidate_images = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                shape_text = shape.text.strip()
                if shape_text:
                    texts.append(shape_text)

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    width_inches = shape.width / 914400
                    height_inches = shape.height / 914400
                    if width_inches >= 1.2 and height_inches >= 0.8:
                        candidate_images.append(
                            (
                                shape.image.blob,
                                shape.image.content_type or "image/png",
                            )
                        )
                except Exception:
                    continue

        slide_text = "\n".join(texts).strip()
        vision_candidate = (
            analyze_visuals
            and bool(candidate_images)
            and (
                len(slide_text) < 500
                or len(candidate_images) >= 2
            )
        )

        units.append({
            "source_location": f"Slide {slide_index}",
            "text": slide_text or "[该页未提取到文字]",
            "source_context": slide_text,
            "visual_evidence": [],
            "image_count": len(candidate_images),
            "visual_analysis_status": (
                "candidate" if vision_candidate else "not_requested"
            ),
            "vision_candidate": vision_candidate,
            "_vision_images": candidate_images[:2],
        })

    return units

def extract_document_units(
    file_bytes,
    filename,
    analyze_visuals=False,
    client=None,
    model_name=None,
    status_callback=None,
):
    extension = Path(filename).suffix.lower()

    if extension == ".txt":
        return make_text_units(decode_txt(file_bytes))

    if extension == ".pdf":
        return extract_pdf_units(
            file_bytes,
            analyze_visuals=analyze_visuals,
            client=client,
            model_name=model_name,
            status_callback=status_callback,
        )

    if extension == ".docx":
        return extract_docx_units(file_bytes)

    if extension == ".pptx":
        return extract_pptx_units(
            file_bytes,
            analyze_visuals=analyze_visuals,
            client=client,
            model_name=model_name,
            status_callback=status_callback,
        )

    raise ValueError(f"暂不支持的文件格式：{extension}")


def visual_text(visual_evidence):
    blocks = []
    for visual in visual_evidence:
        summary = visual.get("summary", "").strip()
        evidence = visual.get("evidence", [])
        if summary:
            blocks.append(f"视觉内容说明：{summary}")
        for item in evidence:
            if isinstance(item, dict) and item.get("text"):
                blocks.append(
                    f"视觉证据（{item.get('kind', 'other')}）：{item['text']}"
                )
    return "\n".join(blocks)


def build_triage_prompt(units):
    rows = []

    for unit in units:
        rows.append({
            "source_location": unit["source_location"],
            "text_preview": compact_text(unit.get("text", ""), 500),
            "image_count": unit.get("image_count", 0),
            "vision_candidate": unit.get("vision_candidate", False),
        })

    return f"""
你是一名投资研究材料筛选助手。目标是节约研究员时间，而不是逐页复述整份材料。

请判断每个材料单元是否值得进入深度分析。

优先保留：
- 公司自身情况、融资、估值、收入、订单、客户、产能、量产
- 产品、核心技术指标、研发进展
- 团队、知识产权、竞争与风险
- 对投资判断有直接帮助的数据和事实

可以降级或跳过：
- 封面、目录、重复页
- 纯行业科普、历史背景、教材型知识
- 与公司无直接关系的宏观叙述
- 纯宣传口号且没有具体事实

宁可保留边界页面，也不要漏掉明显重要的公司事实。

严格返回 JSON 数组，每项：
{{
  "source_location": "Page 1",
  "priority": "core|supporting|skip",
  "reason": "一句话原因",
  "needs_visual": true
}}

needs_visual 只有在文字不足、重要信息主要在图表/截图中时才为 true。

材料索引：
{json.dumps(rows, ensure_ascii=False)}
"""


def canonical_source_location(raw_value, valid_locations):
    raw = str(raw_value or "").strip()

    def key(value):
        value = str(value or "").strip()
        value = value.strip("[](){} \t")
        value = re.sub(r"^UNIT\s+", "", value, flags=re.I)
        value = value.strip("[](){} \t")
        value = re.sub(r"\s*-\s*", "-", value)
        value = re.sub(r"\s+", " ", value)
        return value.lower()

    canonical_map = {
        key(location): location
        for location in valid_locations
    }

    direct = canonical_map.get(key(raw))
    if direct:
        return direct

    match = re.search(
        r"(?i)\b("
        r"Page\s+\d+"
        r"|Slide\s+\d+"
        r"|Paragraphs\s+\d+\s*-\s*\d+"
        r"|Lines\s+\d+\s*-\s*\d+"
        r")\b",
        raw,
    )

    if match:
        return canonical_map.get(key(match.group(1)))

    return None


def source_location_regression_check():
    valid = {
        "Page 14",
        "Slide 3",
        "Paragraphs 1-5",
        "Lines 9-16",
    }
    samples = {
        "Page 14": "Page 14",
        "UNIT Page 14": "Page 14",
        "[UNIT Page 14]": "Page 14",
        "[Page 14]": "Page 14",
        "Slide 3": "Slide 3",
        "[UNIT Slide 3]": "Slide 3",
        "Paragraphs 1 - 5": "Paragraphs 1-5",
        "[UNIT Lines 9-16]": "Lines 9-16",
    }

    for raw, expected in samples.items():
        actual = canonical_source_location(raw, valid)
        if actual != expected:
            raise RuntimeError(
                f"source location normalization failed: "
                f"{raw!r} -> {actual!r}, expected {expected!r}"
            )


HIGH_VALUE_TERMS = (
    "营收", "收入", "利润", "融资", "估值", "合同", "订单", "采购",
    "客户", "专利", "软著", "产能", "量产", "中试线", "销售",
    "团队", "创始人", "董事", "股东", "成本", "毛利", "现金流",
    "revenue", "financing", "valuation", "contract", "order",
    "customer", "patent", "capacity", "production",
)


def has_high_value_signal(text):
    lowered = str(text or "").lower()
    return any(term.lower() in lowered for term in HIGH_VALUE_TERMS)


def triage_units(client, units, model_name):
    if len(units) < TRIAGE_MODEL_MIN_UNITS:
        for unit in units:
            unit["_triage_priority"] = "core"
            unit["_triage_reason"] = (
                "材料单元少于 41 个，直接进入 text-first 分析，"
                "避免筛选调用本身比节省的提取成本更高。"
            )
            unit["_needs_visual_from_triage"] = False
        return units, []

    record_api_call("triage")
    response = client.responses.create(
        model=model_name,
        input=build_triage_prompt(units),
    )
    result = parse_model_json(response.output_text)

    if not isinstance(result, list):
        raise ValueError("页面筛选结果必须是 JSON 数组。")

    valid_locations = {
        unit["source_location"]
        for unit in units
    }
    decision_map = {}

    for row in result:
        if not isinstance(row, dict):
            continue

        location = canonical_source_location(
            row.get("source_location"),
            valid_locations,
        )
        if location:
            decision_map[location] = row

    selected = []
    skipped = []

    for unit in units:
        row = decision_map.get(unit["source_location"], {})
        priority = row.get("priority", "supporting")

        if priority not in {"core", "supporting", "skip"}:
            priority = "supporting"

        unit["_triage_priority"] = priority
        unit["_triage_reason"] = row.get("reason", "")
        unit["_needs_visual_from_triage"] = bool(row.get("needs_visual"))

        if priority == "skip" and has_high_value_signal(unit.get("text", "")):
            priority = "supporting"
            unit["_triage_priority"] = priority
            unit["_triage_reason"] = (
                "本地规则检测到收入/融资/订单/客户/IP/产能等高价值关键词，保守保留。"
            )

        if priority == "skip":
            skipped.append(unit)
        else:
            selected.append(unit)

    # 防止模型过度筛选：至少保留约三分之一材料。
    minimum = max(1, len(units) // 3)
    if len(selected) < minimum:
        selected = units
        skipped = []
        for unit in selected:
            unit["_triage_priority"] = "supporting"

    return selected, skipped


def representative_sample_units(units, limit=5):
    if len(units) <= limit:
        return list(units), []

    positions = []
    for index in range(limit):
        position = round(
            index * (len(units) - 1) / (limit - 1)
        )
        if position not in positions:
            positions.append(position)

    sampled = [
        unit
        for index, unit in enumerate(units)
        if index in positions
    ]
    omitted = [
        unit
        for index, unit in enumerate(units)
        if index not in positions
    ]

    return sampled, omitted


VISUAL_RELATION_TERMS = (
    "竞争格局",
    "客户",
    "股权",
    "组织架构",
    "融资",
    "估值",
    "时间线",
    "路线图",
    "产线",
    "中试线",
)


def needs_visual_followup(unit):
    text = str(unit.get("text", "") or "")
    if unit.get("_needs_visual_from_triage"):
        return True

    if len(text) < 180:
        return True

    relation_page = any(
        term in text
        for term in VISUAL_RELATION_TERMS
    )

    if (
        relation_page
        and len(text) < 900
        and (
            unit.get("drawing_count", 0) >= 10
            or unit.get("image_count", 0) >= 3
        )
    ):
        return True

    if (
        len(text) < 420
        and has_high_value_signal(text)
        and unit.get("image_count", 0) >= 2
    ):
        return True

    return False


def enrich_selected_visuals(
    client,
    units,
    model_name,
    enabled,
    status_callback=None,
):
    if not enabled:
        return 0

    candidates = [
        unit
        for unit in units
        if needs_visual_followup(unit)
        and unit.get("_vision_images")
    ]

    # Vision 只补文字明显不足或关系图/时间线确实影响判断的页面。
    candidates.sort(
        key=lambda item: len(item.get("text", ""))
    )
    candidates = candidates[:MAX_VISION_UNITS]

    analyzed = 0

    for unit in candidates:
        images = unit.get("_vision_images", [])[:1]
        if not images:
            continue

        if status_callback:
            status_callback(
                f"补充读取视觉证据：{unit['source_location']} "
                f"({analyzed + 1}/{len(candidates)})"
            )

        blob, mime_type = images[0]

        try:
            result = analyze_visual(
                client,
                model_name,
                blob,
                mime_type,
                unit["source_location"],
                unit.get("text", ""),
            )
            unit["visual_evidence"] = [result]
            unit["visual_analysis_status"] = result["status"]
            analyzed += 1
        except Exception as exc:
            unit["visual_analysis_status"] = f"failed: {exc}"

    return analyzed


def build_batch_prompt(batch_units):
    blocks = []

    for unit in batch_units:
        visual_description = visual_text(
            unit.get("visual_evidence", [])
        )
        blocks.append(
            f"[UNIT {unit['source_location']}]\n"
            f"[TEXT]\n{compact_text(unit['text'], 6000)}\n"
            f"[VISUAL]\n{visual_description or '无'}"
        )

    material = "\n\n".join(blocks)

    return f"""
你是一名投资研究材料抽取助手。

目标不是“把页面里的所有事实都列出来”，而是提取少量真正会进入研究摘要、研究目录或风险核查的高价值事实。

【优先】
- 收入、利润、融资、估值、订单、采购、客户、合同
- 商业化、量产、产能、产线进展
- 核心产品及关键性能指标
- 关键团队履历
- 专利、软著等知识产权
- 重要竞争定位、风险和限制
- 明确披露的客户/供应商集中度、关键合作方依赖、单一来源依赖
- 监管资质、审批、合规、诉讼、关联交易、股权/治理等可能影响交易判断的重大事项

【不要】
- 单独年份、栏目名、阶段名
- 一般行业科普
- 与公司无直接关系的技术常识
- 同一页把每一个技术参数都拆成独立事实
- 重复表达
- 脱离上下文无法理解的碎片

同一产品的一组相关技术参数，优先合并成一条完整事实。
如果材料只说某技术“可用于金融/医疗/工业”等潜在用途，而没有真实客户、合同、部署或收入证据，必须标为 application_scenario，不能写成已经形成“业务”。
如果材料披露客户名称或匿名客户但没有订单金额，可标为 customer_reference，并保留“公司自述/材料披露”的归属。
attribution 只能来自正文对信息提供者/说话者的明确标注；禁止从文件名、水印、页脚、保密声明、投资机构名称或“private and confidential”等文字推断归属。正文没有明确 speaker 时使用“材料披露”或 Unknown，不要把水印机构写成事实声称者。
每个 UNIT 原则上最多提取 4 条；如果没有真正重要的信息，可以 0 条。

严格返回 JSON 数组。每条包含：
- claim: 可以脱离原文独立读懂的完整事实
- topic: 具体主题
- topic_group: 更高一层的聚合主题，例如“2025收入”“光量子计算机订单”“LNOI光芯片指标”“中试线建设”
- comparison_key: 只有真正可横向比较的同一指标/事项才使用同一个键
- entity: 事实主体，例如“图灵量子”；无法确定填 Unknown
- fact_kind: revenue_actual / revenue_projection / order_signed / purchase_intent / customer_reference / financing / valuation / capacity / production_plan / application_scenario / product / technology_metric / team / ip / market / risk / other
- period: 时间或期间；没有填 Unknown
- category: 只能是 公司概况 / 技术与产品 / 客户与订单 / 商业化与产能 / 财务与预测 / 团队 / 知识产权 / 市场与竞争 / 风险与限制 / 其他
- source_location: 只能复制 UNIT 标签里的位置本身。例如看到 `[UNIT Page 14]` 时，必须输出 `"Page 14"`；禁止输出 `"UNIT Page 14"`、`"[UNIT Page 14]"` 或其他包装
- source_text: 支持该事实的最小必要原始证据，尽量复制原文，不要改写
- attribution: 谁提供或声称该信息；公司自述要保留公司/管理层归属，无法判断填 Unknown
- evidence_status: 只能是 Unverified / Supported / Projection / Subjective / Marketing
- evidence_kind: text / visual / mixed

材料主要为中文时，输出中文。不要替材料增强语义，不要把预测写成已发生事实。

材料：
{material}
"""

def parse_model_json(raw_text):
    raw_text = (raw_text or "").strip()

    if raw_text.startswith("```"):
        lines = raw_text.splitlines()

        if lines and lines[0].startswith("```"):
            lines = lines[1:]

        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]

        raw_text = "\n".join(lines).strip()

    return json.loads(raw_text)


def is_rate_limit_error(exc):
    text = str(exc).lower()
    return (
        "429" in text
        or "rate limit" in text
        or "ratelimit" in text
        or "free-models-per-day" in text
    )


def normalize_claim_schema(item):
    status_aliases = {
        "Marketing": "Subjective / Marketing",
        "Subjective": "Subjective / Marketing",
        "Subjective/Marketing": "Subjective / Marketing",
        "Subjective / Marketing": "Subjective / Marketing",
        "Forecast": "Projection",
        "Projected": "Projection",
        "Projection": "Projection",
        "Verified": "Supported",
        "Supported": "Supported",
        "Unverified": "Unverified",
        "Unknown": "Unverified",
    }

    kind_aliases = {
        "text": "text",
        "Text": "text",
        "visual": "visual",
        "Visual": "visual",
        "image": "visual",
        "Image": "visual",
        "mixed": "mixed",
        "Mixed": "mixed",
        "text+visual": "mixed",
        "text + visual": "mixed",
    }

    if "evidence_status" in item:
        raw_status = str(item["evidence_status"]).strip()
        item["evidence_status"] = status_aliases.get(raw_status, raw_status)

    if "evidence_kind" in item:
        raw_kind = str(item["evidence_kind"]).strip()
        item["evidence_kind"] = kind_aliases.get(raw_kind, raw_kind)

    return item


def validate_claim(item, valid_locations):
    item = normalize_claim_schema(item)

    item.setdefault("topic_group", item.get("topic", "其他"))
    item.setdefault("entity", "Unknown")
    item.setdefault("fact_kind", "other")
    item.setdefault("period", "Unknown")

    for field in (
        "claim",
        "topic",
        "comparison_key",
        "topic_group",
        "entity",
        "fact_kind",
        "period",
        "category",
        "source_location",
        "source_text",
        "attribution",
        "evidence_status",
        "evidence_kind",
    ):
        if field in item and isinstance(item[field], str):
            item[field] = item[field].strip()

    required = {
        "claim",
        "topic",
        "comparison_key",
        "topic_group",
        "entity",
        "fact_kind",
        "period",
        "category",
        "source_location",
        "source_text",
        "attribution",
        "evidence_status",
        "evidence_kind",
    }

    missing = required - item.keys()
    if missing:
        raise ValueError(f"模型输出缺少字段：{sorted(missing)}")

    if item["evidence_status"] not in EVIDENCE_STATUSES:
        raise ValueError(f"不支持的 evidence_status：{item['evidence_status']}")

    if item["evidence_kind"] not in {"text", "visual", "mixed"}:
        raise ValueError(f"不支持的 evidence_kind：{item['evidence_kind']}")

    if item["category"] not in CATEGORIES:
        item["category"] = "其他"

    matched_location = canonical_source_location(
        item["source_location"],
        valid_locations,
    )

    if matched_location:
        item["source_location"] = matched_location
    else:
        raise ValueError(
            f"source_location 不在当前批次：{item['source_location']}"
        )

    if len(str(item["claim"]).strip()) < 8:
        raise ValueError(f"claim 过短，无法独立理解：{item['claim']}")

    return item


def context_snippet(unit_text, source_text, radius=260):
    unit_text = (unit_text or "").strip()
    source_text = (source_text or "").strip()

    if not unit_text:
        return ""

    index = unit_text.find(source_text)

    if index >= 0 and source_text:
        start = max(0, index - radius)
        end = min(len(unit_text), index + len(source_text) + radius)
        snippet = unit_text[start:end].strip()

        if start > 0:
            snippet = "…" + snippet

        if end < len(unit_text):
            snippet += "…"

        return snippet

    return compact_text(unit_text, 650)


def canonicalize_dates(text):
    text = unicodedata.normalize(
        "NFKC",
        str(text or ""),
    )

    def repl_cn(match):
        year = int(match.group(1))
        month = int(match.group(2))
        day = match.group(3)
        if day:
            return f"{year:04d}.{month:02d}.{int(day):02d}"
        return f"{year:04d}.{month:02d}"

    text = re.sub(
        r"(20\d{2})\s*年\s*(\d{1,2})\s*月"
        r"(?:\s*(\d{1,2})\s*[日号])?",
        repl_cn,
        text,
    )

    def repl_sep(match):
        year = int(match.group(1))
        month = int(match.group(2))
        day = match.group(3)
        if day:
            return f"{year:04d}.{month:02d}.{int(day):02d}"
        return f"{year:04d}.{month:02d}"

    text = re.sub(
        r"(?<!\d)(20\d{2})[./\-](\d{1,2})"
        r"(?:[./\-](\d{1,2}))?(?!\d)",
        repl_sep,
        text,
    )
    return text


def normalize_for_match(text):
    text = canonicalize_dates(text).lower()
    chars = []
    for char in text:
        category = unicodedata.category(char)
        if char.isspace():
            continue
        if category.startswith("P") or category.startswith("Z"):
            continue
        chars.append(char)
    return "".join(chars)


def numeric_tokens(text):
    text = canonicalize_dates(text)
    dates = re.findall(
        r"(?<!\d)20\d{2}\.\d{2}(?:\.\d{2})?(?!\d)",
        text,
    )
    text_without_dates = re.sub(
        r"(?<!\d)20\d{2}\.\d{2}(?:\.\d{2})?(?!\d)",
        " ",
        text,
    )
    numbers = re.findall(
        r"\d+(?:\.\d+)?",
        text_without_dates,
    )
    return dates + numbers


def ordered_subsequence_ratio(needle, haystack):
    if not needle:
        return 1.0

    matched = 0
    for char in haystack:
        if matched < len(needle) and char == needle[matched]:
            matched += 1

    return matched / len(needle)


def ngram_anchor_coverage(needle, haystack, width=8):
    if not needle:
        return 1.0
    if len(needle) < width:
        return 1.0 if needle in haystack else 0.0

    windows = [
        needle[index:index + width]
        for index in range(len(needle) - width + 1)
    ]
    matched = sum(
        window in haystack
        for window in windows
    )
    return matched / len(windows)


def source_anchor_matches(source_text, unit_text):
    normalized_source = normalize_for_match(
        source_text
    )
    normalized_unit = normalize_for_match(
        unit_text
    )

    if not normalized_source:
        return True
    if normalized_source in normalized_unit:
        return True

    raw_segments = re.split(
        r"[；;。！？!?，,\n]+",
        str(source_text or ""),
    )
    segments = [
        normalize_for_match(segment)
        for segment in raw_segments
        if len(normalize_for_match(segment)) >= 4
    ]

    if not segments:
        return False

    matched = [
        segment
        for segment in segments
        if segment in normalized_unit
    ]
    matched_chars = sum(len(segment) for segment in matched)
    total_chars = sum(len(segment) for segment in segments)

    if total_chars and matched_chars / total_chars >= 0.65:
        return True

    if (
        len(segments) >= 2
        and len(matched) >= min(2, len(segments))
    ):
        return True

    # PDF pages often interleave headings, bullets and chart labels between
    # evidence fragments. Models may therefore return a source_text made of
    # several correctly ordered page fragments that are not one contiguous
    # substring. Accept that case only when both local phrase coverage and
    # global ordering are strong; numeric support is still checked separately.
    if len(normalized_source) >= 24:
        local_coverage = ngram_anchor_coverage(
            normalized_source,
            normalized_unit,
            width=8,
        )
        ordered_coverage = ordered_subsequence_ratio(
            normalized_source,
            normalized_unit,
        )
        if (
            local_coverage >= 0.50
            and ordered_coverage >= 0.90
        ):
            return True

    return False


def normalize_attribution_value(value, source_text):
    value = str(value or "").strip()
    if not value:
        return "Unknown"

    if value.endswith("材料披露"):
        prefix = value[:-4].strip(" /、·")
        if (
            prefix
            and normalize_for_match(prefix)
            not in normalize_for_match(source_text)
        ):
            return "材料披露"

    return value


def evidence_check(claim, source_text, unit, evidence_kind):
    claim_numbers = set(numeric_tokens(claim))
    source_numbers = set(numeric_tokens(source_text))
    missing_numbers = sorted(
        claim_numbers - source_numbers
    )

    source_anchor_found = (
        source_anchor_matches(
            source_text,
            unit.get("text", ""),
        )
        or evidence_kind in {"visual", "mixed"}
    )

    issues = []

    if missing_numbers:
        issues.append(
            "claim 中出现了 source_text 未直接支持的数字："
            + ", ".join(missing_numbers)
        )

    if not source_anchor_found:
        issues.append(
            "source_text 与该材料单元的文字锚点匹配不足"
        )

    return {
        "status": "pass" if not issues else "review",
        "issues": issues,
    }

def extract_claims_from_batch(
    client,
    batch_units,
    document_name,
    source_type,
    model_name,
):
    record_api_call("claim")
    response = client.responses.create(
        model=model_name,
        input=build_batch_prompt(batch_units),
    )

    extracted = parse_model_json(response.output_text)

    if not isinstance(extracted, list):
        raise ValueError("模型输出必须是 JSON 数组。")

    unit_map = {
        unit["source_location"]: unit
        for unit in batch_units
    }
    valid_locations = set(unit_map)
    claims = []
    item_warnings = []

    for item_index, item in enumerate(extracted, start=1):
        try:
            item = validate_claim(
                item,
                valid_locations,
            )
        except Exception as exc:
            item_warnings.append(
                f"第 {item_index} 条模型输出已跳过：{exc}"
            )
            continue

        unit = unit_map[item["source_location"]]
        item["attribution"] = normalize_attribution_value(
            item.get("attribution", "Unknown"),
            item.get("source_text", ""),
        )

        claims.append({
            "claim": item["claim"],
            "topic": item["topic"],
            "topic_group": item["topic_group"],
            "comparison_key": item["comparison_key"],
            "entity": item["entity"],
            "fact_kind": item["fact_kind"],
            "period": item["period"],
            "category": item["category"],
            "source_text": item["source_text"],
            "source_context": context_snippet(
                unit.get("text", ""),
                item["source_text"],
            ),
            "attribution": item["attribution"],
            "evidence_status": item["evidence_status"],
            "evidence_kind": item["evidence_kind"],
            "document_id": Path(document_name).stem,
            "source_filename": document_name,
            "source_type": source_type,
            "source_location": item["source_location"],
            "image_count": unit.get("image_count", 0),
            "visual_evidence": unit.get("visual_evidence", []),
            "visual_analysis_status": unit.get(
                "visual_analysis_status",
                "not_applicable",
            ),
            "_preview_bytes": unit.get(
                "preview_bytes"
            ),
        })

        claims[-1]["evidence_check"] = evidence_check(
            claims[-1]["claim"],
            claims[-1]["source_text"],
            unit,
            claims[-1]["evidence_kind"],
        )

    return claims, item_warnings

def demo_topic(claim):
    text = claim["claim"]

    if "聚焦工商业储能" in text or "能源管理服务" in text:
        return "业务概览"

    if "管理团队" in text or "CEO" in text or "CTO" in text:
        return "核心团队"

    if "员工" in text or "研发人员" in text or "博士" in text:
        return "人员结构"

    if "项目" in text:
        return "累计完成项目数"

    if "营收" in text or "3.2亿元" in text:
        return "2026年营收预测"

    if "专利" in text:
        return "核心专利情况"

    if "领先" in text or "定位" in text:
        return "市场定位"

    return "其他"


def demo_category(claim):
    text = claim["claim"]

    if "聚焦工商业储能" in text or "能源管理服务" in text:
        return "公司概况"

    if (
        "管理团队" in text
        or "CEO" in text
        or "CTO" in text
        or "员工" in text
        or "研发人员" in text
        or "博士" in text
    ):
        return "团队"

    if "项目" in text:
        return "商业化与产能"

    if "营收" in text or "3.2亿元" in text:
        return "财务与预测"

    if "专利" in text:
        return "知识产权"

    if "领先" in text or "定位" in text:
        return "市场与竞争"

    return "其他"


def normalize_demo_claims(claims):
    for index, item in enumerate(claims, start=1):
        item["claim_id"] = f"claim_{index}"
        item["topic"] = demo_topic(item)
        item["comparison_key"] = item["topic"]
        item["category"] = demo_category(item)
        item.setdefault("source_filename", f"{item['document_id']}.txt")
        item.setdefault("source_type", "txt")
        start = item.get("source_line_start", "?")
        end = item.get("source_line_end", "?")
        item.setdefault("source_location", f"Lines {start}-{end}")
        item.setdefault("source_context", item.get("source_text", ""))
        item.setdefault("evidence_kind", "text")
        item.setdefault("image_count", 0)
        item.setdefault("visual_evidence", [])
        item.setdefault("visual_analysis_status", "not_applicable")
        item.setdefault("_preview_bytes", None)
        item.setdefault("topic_group", item["topic"])
        item.setdefault("entity", "星云能源科技有限公司")

        if "营收" in item["claim"] or "3.2亿元" in item["claim"]:
            item.setdefault("fact_kind", "revenue_projection")
            item.setdefault("period", "2026")
        elif "专利" in item["claim"]:
            item.setdefault("fact_kind", "ip")
            item.setdefault("period", "Unknown")
        elif (
            "管理团队" in item["claim"]
            or "CEO" in item["claim"]
            or "CTO" in item["claim"]
            or "员工" in item["claim"]
            or "研发人员" in item["claim"]
            or "博士" in item["claim"]
        ):
            item.setdefault("fact_kind", "team")
            item.setdefault("period", "2026")
        elif "项目" in item["claim"]:
            item.setdefault("fact_kind", "other")
            item.setdefault("period", "2026")
        elif "领先" in item["claim"] or "定位" in item["claim"]:
            item.setdefault("fact_kind", "market")
            item.setdefault("period", "Unknown")
        else:
            item.setdefault("fact_kind", "other")
            item.setdefault("period", "Unknown")

        item.setdefault("evidence_check", {"status": "pass", "issues": []})



def build_demo_research_output(claims):
    def ids_containing(text):
        return [
            item["claim_id"]
            for item in claims
            if text in item["claim"]
        ]

    project_ids = [
        item["claim_id"]
        for item in claims
        if "工商业储能项目" in item["claim"]
    ]
    revenue_ids = [
        item["claim_id"]
        for item in claims
        if "营收" in item["claim"]
    ]
    patent_ids = [
        item["claim_id"]
        for item in claims
        if "专利" in item["claim"]
    ]

    sections = fallback_research_output(
        claims,
        output_language=current_ui_language(),
    )["sections"]

    # Demo regression targets should exercise the same synthesis semantics as Live:
    # preserve material caveats and resolve explicit forecast revisions to current state.
    for section in sections:
        for group in section.get("groups", []):
            if group.get("title") == "累计完成项目数":
                group["summary"] = (
                    "截至2026年6月累计完成30个工商业储能项目，"
                    "较2026年3月的27个继续增加。"
                )
                group["fact_ids"] = project_ids
            elif group.get("title") == "2026年营收预测":
                group["summary"] = (
                    "2026年营收目标已由3亿元调整至3.2亿元，"
                    "尚未形成已实现收入。"
                )
                group["fact_ids"] = revenue_ids
            elif group.get("title") == "权属与合作开发":
                group["summary"] = (
                    "公司称11项核心专利中有5项为合作开发；"
                    "现有材料未说明这5项与申请/开发阶段专利是否重叠。"
                )
                group["fact_ids"] = patent_ids
            elif group.get("title") == "授权状态":
                group["summary"] = (
                    "3项仍处于申请或开发阶段；现有材料不足以据此确认"
                    "公司独立持有且已授权的核心专利数量。"
                )
                group["fact_ids"] = patent_ids

    return {
        "headline": "",
        "executive_summary": [
            {
                "text": "公司材料显示，项目数量从2026年3月的27个增加至6月的30个，属于不同时间点的业务进展，不应直接视为冲突。",
                "fact_ids": project_ids,
            },
            {
                "text": "管理层将2026年全年营收目标从3亿元调整为3.2亿元，3.2亿元仍属于内部经营目标而非已实现收入。",
                "fact_ids": revenue_ids,
            },
            {
                "text": "公司称涉及11项核心专利，但其中包含合作开发以及仍处于申请或开发阶段的项目。",
                "fact_ids": patent_ids,
            },
        ],
        "key_metrics": [
            {
                "label": "最新项目数",
                "value": "30个",
                "note": "截至2026年6月，公司材料",
                "fact_ids": ids_containing("30个工商业储能项目"),
            },
            {
                "label": "2026营收目标",
                "value": "3.2亿元",
                "note": "Projection / 内部经营目标",
                "fact_ids": ids_containing("3.2亿元"),
            },
            {
                "label": "核心专利口径",
                "value": "11项",
                "note": "公司称，非全部已授权",
                "fact_ids": patent_ids,
            },
        ],
        "sections": sections,
        "research_flags": [
            {
                "title": "核心专利权属与授权状态未完整拆分",
                "type": "material_evidence_gap",
                "unresolved_question": (
                    "11项核心专利中，有多少项为公司独立持有且已经授权？"
                    "5项合作开发与3项申请/开发阶段专利是否存在重叠？"
                ),
                "why": (
                    "现有材料给出了总量、合作开发数量和未授权/未完成数量，"
                    "但没有披露这些分类之间的集合关系，因此无法确认公司实际可控制的已授权核心IP数量。"
                ),
                "action": (
                    "获取完整专利清单，核对专利号、申请人/专利权人、法律状态、"
                    "共同权利人以及合作开发或许可安排。"
                ),
                "fact_ids": patent_ids,
            }
        ],
    }


def localize_demo_english(claims, research_output):
    claim_translations = {
        "公司表示，截至2026年3月已累计完成27个工商业储能项目。": "The company stated that it had completed 27 commercial and industrial energy-storage projects as of March 2026.",
        "管理层预计2026年全年营收达到3亿元。": "Management projected full-year 2026 revenue of RMB 3 hundred million.",
        "截至2026年6月，公司累计完成30个工商业储能项目。": "As of June 2026, the company had completed 30 commercial and industrial energy-storage projects.",
        "公司将自己定位为国内领先的工商业储能服务商。": "The company positions itself as a leading domestic commercial and industrial energy-storage service provider.",
        "管理层将2026年全年营收目标调整为3.2亿元。": "Management revised its full-year 2026 revenue target to RMB 3.2 hundred million.",
        "3.2亿元为内部经营目标，尚未形成已实现收入。": "RMB 3.2 hundred million is an internal operating target, not realized revenue.",
        "公司称目前涉及11项核心专利，其中5项为合作开发。": "The company states that 11 core patents are involved, of which 5 are co-developed.",
        "其中3项仍处于申请或开发阶段，尚未形成已授权专利。": "Of these, 3 remain at the application or development stage and have not become granted patents.",
        "公司材料称，公司聚焦工商业储能系统集成与能源管理服务。": "Company materials state that the business focuses on commercial and industrial energy-storage system integration and energy-management services.",
        "公司现有员工86人，其中研发人员52人。": "The company has 86 employees, including 52 R&D staff.",
        "公司称研发团队中有12名博士。": "The company states that its R&D team includes 12 PhD holders.",
        "公司核心管理团队包括CEO、CTO与销售负责人。": "The core management team includes the CEO, CTO and head of sales.",
    }
    attribution_translations = {
        "公司/管理层访谈": "Company / management interview",
        "管理层": "Management",
        "公司投资人材料": "Company investor materials",
        "公司": "Company",
        "公司更新预算": "Company updated budget",
        "公司知识产权说明": "Company IP note",
        "公司材料": "Company materials",
    }
    text_translations = {
        "截至2026年6月累计完成30个工商业储能项目，较2026年3月的27个继续增加。": "As of June 2026, 30 commercial and industrial energy-storage projects had been completed, up from 27 in March 2026.",
        "2026年营收目标已由3亿元调整至3.2亿元，尚未形成已实现收入。": "The 2026 revenue target was revised from RMB 3 hundred million to RMB 3.2 hundred million; this remains a target rather than realized revenue.",
        "公司称11项核心专利中有5项为合作开发；现有材料未说明这5项与申请/开发阶段专利是否重叠。": "The company states that 5 of 11 core patents are co-developed; the supplied materials do not clarify whether these overlap with patents still at the application/development stage.",
        "3项仍处于申请或开发阶段；现有材料不足以据此确认公司独立持有且已授权的核心专利数量。": "Three remain at the application or development stage; the supplied materials are insufficient to determine how many core patents are both independently controlled by the company and already granted.",
        "公司材料显示，项目数量从2026年3月的27个增加至6月的30个，属于不同时间点的业务进展，不应直接视为冲突。": "Company materials show project count increasing from 27 in March 2026 to 30 in June 2026. These are different time points and should not automatically be treated as a conflict.",
        "管理层将2026年全年营收目标从3亿元调整为3.2亿元，3.2亿元仍属于内部经营目标而非已实现收入。": "Management revised the full-year 2026 revenue target from RMB 3 hundred million to RMB 3.2 hundred million; the latter remains an internal operating target rather than realized revenue.",
        "公司称涉及11项核心专利，但其中包含合作开发以及仍处于申请或开发阶段的项目。": "The company states that 11 core patents are involved, but this set includes co-developed items and items still at the application or development stage.",
        "核心专利权属与授权状态未完整拆分": "Core patent ownership and grant status are not fully separated",
        "11项核心专利中，有多少项为公司独立持有且已经授权？5项合作开发与3项申请/开发阶段专利是否存在重叠？": "How many of the 11 core patents are independently held by the company and already granted? Do the 5 co-developed patents overlap with the 3 patents at the application/development stage?",
        "现有材料给出了总量、合作开发数量和未授权/未完成数量，但没有披露这些分类之间的集合关系，因此无法确认公司实际可控制的已授权核心IP数量。": "The materials provide the total count, co-developed count and not-yet-granted/development-stage count, but do not disclose how these sets overlap, so the number of granted core IP assets actually controlled by the company cannot be confirmed.",
        "获取完整专利清单，核对专利号、申请人/专利权人、法律状态、共同权利人以及合作开发或许可安排。": "Obtain the complete patent list and verify patent numbers, applicants/owners, legal status, co-owners and co-development or licensing arrangements.",
    }
    title_translations = {
        "业务概览": "Business overview",
        "核心团队": "Core team",
        "人员结构": "Team structure",
        "累计完成项目数": "Completed project count",
        "2026年营收预测": "2026 revenue forecast",
        "核心专利情况": "Core patent position",
        "市场定位": "Market positioning",
        "授权状态": "Grant status",
        "权属与合作开发": "Ownership & co-development",
        "专利组合": "Patent portfolio",
        "最新项目数": "Latest project count",
        "2026营收目标": "2026 revenue target",
        "核心专利口径": "Core patent scope",
    }
    note_translations = {
        "截至2026年6月，公司材料": "As of June 2026; company materials",
        "Projection / 内部经营目标": "Projection / internal operating target",
        "公司称，非全部已授权": "Company-reported; not all are granted",
    }

    for item in claims:
        item["claim"] = claim_translations.get(item.get("claim", ""), item.get("claim", ""))
        item["attribution"] = attribution_translations.get(item.get("attribution", ""), item.get("attribution", ""))

    for item in research_output.get("executive_summary", []):
        item["text"] = text_translations.get(item.get("text", ""), item.get("text", ""))

    for metric in research_output.get("key_metrics", []):
        metric["label"] = title_translations.get(metric.get("label", ""), metric.get("label", ""))
        metric["note"] = note_translations.get(metric.get("note", ""), metric.get("note", ""))
        if metric.get("value") == "30个":
            metric["value"] = "30"
        elif metric.get("value") == "3.2亿元":
            metric["value"] = "RMB 3.2 hundred million"
        elif metric.get("value") == "11项":
            metric["value"] = "11"

    for section in research_output.get("sections", []):
        for group in section.get("groups", []):
            group["title"] = title_translations.get(group.get("title", ""), group.get("title", ""))
            group["summary"] = text_translations.get(
                group.get("summary", ""),
                claim_translations.get(group.get("summary", ""), group.get("summary", "")),
            )

    for flag in research_output.get("research_flags", []):
        for key in ("title", "unresolved_question", "why", "action"):
            flag[key] = text_translations.get(flag.get(key, ""), flag.get(key, ""))

    return claims, research_output



def add_claim_ids(claims):
    for index, item in enumerate(claims, start=1):
        item["claim_id"] = f"claim_{index}"




def compact_fact_for_model(item):
    return {
        "fact_id": item["claim_id"],
        "claim": item["claim"],
        "category": item.get("category", "其他"),
        "topic_group": item.get("topic_group", item.get("topic", "其他")),
        "comparison_key": item.get("comparison_key", ""),
        "entity": item.get("entity", "Unknown"),
        "fact_kind": item.get("fact_kind", "other"),
        "period": item.get("period", "Unknown"),
        "attribution": item.get("attribution", "Unknown"),
        "evidence_status": item.get("evidence_status", "Unverified"),
        "evidence_check": item.get("evidence_check", {}).get("status", "review"),
        "source": (
            f"{item.get('source_filename', '')} · "
            f"{item.get('source_location', '')}"
        ),
        "source_text": compact_text(item.get("source_text", ""), 320),
    }


def canonical_matrix_domain(section_title, items):
    title = str(section_title or "").strip()
    if title in MATRIX_DOMAIN_ORDER:
        return title

    categories = {
        item.get("category", "其他")
        for item in (items or [])
    }
    kinds = {
        item.get("fact_kind", "other")
        for item in (items or [])
    }

    if "公司概况" in categories or any(
        keyword in title
        for keyword in ("公司概况", "公司简介", "业务概览")
    ):
        return "公司概览"

    if (
        categories
        & {"客户与订单", "商业化与产能", "市场与竞争"}
        or kinds
        & {
            "order_signed",
            "purchase_intent",
            "customer_reference",
            "capacity",
            "production_plan",
            "market",
        }
        or any(
            keyword in title
            for keyword in (
                "客户",
                "订单",
                "采购",
                "商业化",
                "产能",
                "量产",
                "市场",
                "竞争",
            )
        )
    ):
        return "商业化与竞争"

    if (
        "财务与预测" in categories
        or kinds
        & {
            "revenue_actual",
            "revenue_projection",
            "financing",
            "valuation",
        }
        or any(
            keyword in title
            for keyword in ("财务", "收入", "营收", "融资", "估值")
        )
    ):
        return "财务与融资"

    if (
        "团队" in categories
        or "team" in kinds
        or any(
            keyword in title
            for keyword in ("团队", "创始", "管理层", "人员", "员工")
        )
    ):
        return "团队与组织"

    if (
        "技术与产品" in categories
        or kinds
        & {
            "technology_metric",
            "product",
            "application_scenario",
        }
        or any(
            keyword in title
            for keyword in ("技术", "产品", "应用场景", "解决方案")
        )
    ):
        return "产品与技术"

    if (
        "知识产权" in categories
        or "ip" in kinds
        or any(keyword in title for keyword in ("知识产权", "专利", "软著"))
    ):
        return "知识产权"

    return "其他重要事项"


def canonical_matrix_topic(label, items, domain):
    label = str(label or "").strip() or "其他"
    text = " ".join(
        [label]
        + [str(item.get("claim", "")) for item in (items or [])]
    )

    if domain == "知识产权" and label in {
        "核心专利情况",
        "专利情况",
        "知识产权",
        "专利",
    }:
        has_auth = any(
            keyword in text
            for keyword in ("授权", "申请", "开发阶段")
        )
        has_rights = any(
            keyword in text
            for keyword in ("合作开发", "权属", "许可", "独占")
        )
        if has_auth and not has_rights:
            return "授权状态"
        if has_rights and not has_auth:
            return "权属与合作开发"
        return "专利组合"

    if domain == "团队与组织":
        if label in {"团队", "团队情况", "人员情况"}:
            if any(
                keyword in text
                for keyword in ("员工", "研发人员", "博士", "占比", "人员结构")
            ):
                return "人员结构"
            return "核心团队"

    if domain == "商业化与竞争":
        if label in {"客户情况", "客户进展", "客户合作"}:
            return "客户与合作"
        if label in {"订单情况", "合同情况", "订单进展"}:
            return "订单与合同"
        if label in {"产能情况", "量产进展", "产线进展"}:
            return "产能与量产"

    return label


def matrix_topic_priority(label, items, domain):
    text = " ".join(
        [str(label or "")]
        + [str(item.get("claim", "")) for item in (items or [])]
    )

    keyword_orders = {
        "公司概览": [
            ("业务概览", "主营", "业务"),
            ("发展阶段", "成立", "阶段"),
        ],
        "商业化与竞争": [
            ("市场定位", "市场空间", "行业定位"),
            ("竞争格局", "竞争", "竞品"),
            ("客户", "合作客户"),
            ("订单", "合同", "采购"),
            ("项目", "交付", "落地"),
            ("商业化", "变现"),
            ("产能", "量产", "中试", "产线"),
        ],
        "财务与融资": [
            ("历史收入", "已实现收入", "营收"),
            ("预测", "目标", "预算"),
            ("融资",),
            ("估值",),
            ("利润", "成本", "现金流"),
        ],
        "团队与组织": [
            ("创始", "核心团队", "管理团队", "CEO", "CTO"),
            ("员工", "人员规模"),
            ("人员结构", "研发人员", "博士", "占比"),
            ("研发团队", "研发能力"),
            ("销售团队", "销售人员"),
        ],
        "产品与技术": [
            ("核心产品", "产品"),
            ("技术路线", "架构", "方案"),
            ("性能", "指标", "参数"),
            ("应用场景", "潜在用途"),
        ],
        "知识产权": [
            ("专利组合", "专利总览", "核心专利"),
            ("授权状态", "已授权", "申请"),
            ("权属", "合作开发", "许可", "独占"),
            ("技术覆盖", "覆盖"),
            ("风险", "纠纷"),
        ],
    }

    for index, keywords in enumerate(keyword_orders.get(domain, [])):
        if any(keyword in text for keyword in keywords):
            return index

    return 99


def fact_recency_key(item):
    text = (
        str(item.get("claim", ""))
        + " "
        + str(item.get("source_text", ""))
    )

    full_dates = re.findall(
        r"(20\d{2})[年\-/\.](\d{1,2})[月\-/\.]?(\d{1,2})?",
        text,
    )
    if full_dates:
        year, month, day = full_dates[-1]
        return int(year) * 10000 + int(month) * 100 + int(day or 0)

    years = re.findall(r"(20\d{2})年?", text)
    if years:
        return int(years[-1]) * 10000

    return 0


def fallback_research_output(claims, output_language="zh-CN"):
    eligible = [
        item for item in claims
        if item.get("evidence_check", {}).get("status") == "pass"
    ]
    if not eligible:
        eligible = claims

    by_domain = defaultdict(list)
    for item in eligible:
        domain = canonical_matrix_domain(
            item.get("category", "其他"),
            [item],
        )
        by_domain[domain].append(item)

    sections = []
    for domain in MATRIX_DOMAIN_ORDER:
        items = by_domain.get(domain, [])
        if not items:
            continue

        topic_map = defaultdict(list)
        for item in items:
            raw_topic = (
                item.get("topic_group")
                or item.get("topic")
                or "其他"
            )
            topic = canonical_matrix_topic(
                raw_topic,
                [item],
                domain,
            )
            topic_map[topic].append(item)

        groups = []
        for title, group_items in topic_map.items():
            indexed = list(enumerate(group_items))
            _, latest_item = max(
                indexed,
                key=lambda pair: (
                    fact_recency_key(pair[1]),
                    pair[0],
                ),
            )

            groups.append({
                "title": title,
                "summary": latest_item["claim"],
                "fact_ids": [
                    item["claim_id"]
                    for item in group_items[:8]
                ],
            })

        groups.sort(
            key=lambda group: matrix_topic_priority(
                group["title"],
                [
                    item
                    for item in items
                    if item["claim_id"] in group["fact_ids"]
                ],
                domain,
            )
        )

        sections.append({
            "title": domain,
            "summary": (
                f"{len(items)} high-value facts organized."
                if output_language == "en"
                else f"共整理 {len(items)} 条高价值事实。"
            ),
            "groups": groups[:8],
        })

    # Local fallback deliberately does not create Research Flags.
    # Evidence-location warnings are processing/traceability issues,
    # not investment diligence questions under the
    # Material + Unresolved + Actionable contract.
    flags = []

    return {
        "headline": (
            "Material compression complete; every item below remains traceable to source evidence."
            if output_language == "en"
            else "已完成材料压缩；以下内容均可回到原始来源核对。"
        ),
        "executive_summary": [
            {
                "text": item["claim"],
                "fact_ids": [item["claim_id"]],
            }
            for item in eligible[:5]
        ],
        "key_metrics": [],
        "sections": sections,
        "research_flags": flags,
    }

def metric_snapshot_score(metric, claim_map):
    items = [
        claim_map[fact_id]
        for fact_id in metric.get("fact_ids", [])
        if fact_id in claim_map
    ]
    kinds = {
        item.get("fact_kind", "other")
        for item in items
    }
    text = (
        str(metric.get("label", ""))
        + " "
        + str(metric.get("note", ""))
    )

    if kinds & {
        "revenue_actual",
        "revenue_projection",
        "financing",
        "valuation",
        "order_signed",
        "purchase_intent",
        "capacity",
        "production_plan",
    }:
        return 100

    if "ip" in kinds:
        return 80

    score = 20
    if any(
        keyword in text
        for keyword in (
            "集成",
            "线路",
            "光源",
            "迭代",
            "产能",
            "吞吐",
            "良率",
            "周期",
        )
    ):
        score += 20
    if any(
        keyword in text
        for keyword in (
            "层数",
            "模式数",
            "误差",
        )
    ):
        score -= 5
    return score


def prune_key_metrics(metrics, claim_map):
    best_by_fact_set = {}

    for index, metric in enumerate(metrics):
        signature = tuple(sorted(
            fact_id
            for fact_id in metric.get(
                "fact_ids",
                [],
            )
            if fact_id in claim_map
        ))
        if not signature:
            continue

        scored = (
            metric_snapshot_score(
                metric,
                claim_map,
            ),
            -index,
            metric,
        )
        current = best_by_fact_set.get(
            signature
        )
        if (
            current is None
            or scored[:2] > current[:2]
        ):
            best_by_fact_set[signature] = scored

    ranked = sorted(
        best_by_fact_set.values(),
        key=lambda item: (
            -item[0],
            -item[1],
        ),
    )

    selected = []
    tech_count = 0

    for _, _, metric in ranked:
        items = [
            claim_map[fact_id]
            for fact_id in metric.get(
                "fact_ids",
                [],
            )
            if fact_id in claim_map
        ]
        kinds = {
            item.get("fact_kind", "other")
            for item in items
        }
        is_technical = bool(
            kinds & {
                "technology_metric",
                "product",
                "application_scenario",
            }
        )

        if is_technical and tech_count >= 2:
            continue
        if (
            is_technical
            and metric_snapshot_score(
                metric,
                claim_map,
            ) < 20
        ):
            continue

        selected.append(metric)
        if is_technical:
            tech_count += 1
        if len(selected) >= 6:
            break

    return selected


def apply_partial_scope_flag_guard(
    research_output,
    partial_scope,
):
    if not partial_scope:
        return research_output, 0

    structural_terms = (
        "重叠",
        "集合",
        "互斥",
        "对应关系",
        "口径",
        "定义范围",
        "定义",
        "矛盾",
        "冲突",
    )

    kept = []
    dropped = 0

    for flag in research_output.get(
        "research_flags",
        [],
    ):
        flag_type = str(
            flag.get("type", "")
        )
        text = " ".join(
            [
                str(flag.get("title", "")),
                str(flag.get(
                    "unresolved_question",
                    "",
                )),
                str(flag.get("why", "")),
            ]
        )

        if flag_type != "material_evidence_gap":
            kept.append(flag)
            continue

        if any(
            term in text
            for term in structural_terms
        ):
            kept.append(flag)
            continue

        dropped += 1

    research_output = json.loads(
        json.dumps(
            research_output,
            ensure_ascii=False,
        )
    )
    research_output["research_flags"] = kept
    return research_output, dropped


def filter_matrix_primary_home_sections(sections, claim_map):
    filtered_sections = []

    for section in sections:
        section_domain = str(section.get("title", "")).strip()
        groups = []

        for group in section.get("groups", []):
            items = [
                claim_map[fact_id]
                for fact_id in group.get("fact_ids", [])
                if fact_id in claim_map
            ]
            if not items:
                continue

            primary_domains = {
                canonical_matrix_domain(
                    item.get("category", "其他"),
                    [item],
                )
                for item in items
            }

            # Keep a topic where at least one underlying fact naturally belongs
            # to this domain. This removes "catch-all" repetitions such as
            # technology coverage copied into IP or generic risk summaries
            # copied into Other, while preserving mixed topics with a real home.
            if section_domain in MATRIX_DOMAIN_ORDER:
                if section_domain not in primary_domains:
                    continue

            groups.append(group)

        if groups:
            filtered_sections.append({
                **section,
                "groups": groups,
            })

    return filtered_sections


MATERIAL_CAVEAT_TERMS = (
    "需核对",
    "口径不明",
    "口径待核对",
    "待确认",
    "无法确认",
    "未说明对应",
    "未说明口径",
    "needs reconciliation",
    "requires reconciliation",
    "scope is unclear",
    "definition is unclear",
    "to be confirmed",
    "cannot confirm",
    "not specified",
    "remains unclear",
)

MATERIAL_FLAG_KINDS = {
    "revenue_actual",
    "revenue_projection",
    "financing",
    "valuation",
    "order_signed",
    "purchase_intent",
    "customer_reference",
    "capacity",
    "production_plan",
    "ip",
}


def ensure_material_caveat_flags(
    research_output,
    claims,
    partial_scope=False,
    output_language="zh-CN",
):
    if partial_scope:
        return research_output, 0

    claim_map = {
        item["claim_id"]: item
        for item in claims
    }
    output = json.loads(
        json.dumps(
            research_output,
            ensure_ascii=False,
        )
    )
    flags = output.setdefault("research_flags", [])

    existing_ids = [
        set(flag.get("fact_ids", []))
        for flag in flags
    ]

    candidates = []

    for section in output.get("sections", []):
        for group in section.get("groups", []):
            candidates.append((
                str(group.get("title", "")).strip() or "研究主题",
                str(group.get("summary", "")).strip(),
                list(group.get("fact_ids", [])),
            ))

    for item in output.get("executive_summary", []):
        candidates.append((
            "核心摘要中的未解决口径",
            str(item.get("text", "")).strip(),
            list(item.get("fact_ids", [])),
        ))

    added = 0

    for title, text, fact_ids in candidates:
        if not text or not any(
            term.lower() in text.lower()
            for term in MATERIAL_CAVEAT_TERMS
        ):
            continue

        ids = [
            fact_id
            for fact_id in fact_ids
            if fact_id in claim_map
        ]
        if not ids:
            continue

        items = [claim_map[fact_id] for fact_id in ids]
        kinds = {
            item.get("fact_kind", "other")
            for item in items
        }
        if not (kinds & MATERIAL_FLAG_KINDS):
            continue

        current_set = set(ids)
        if any(current_set & existing for existing in existing_ids):
            continue

        is_english = output_language == "en"

        if kinds & {"valuation", "financing"}:
            why = (
                "This scope directly affects financing terms, valuation comparison and dilution analysis."
                if is_english
                else "该口径会直接影响融资条件、估值比较与稀释判断。"
            )
            action = (
                "Reconcile the relevant round, currency, pre-/post-money basis, financing amount and key terms."
                if is_english
                else "核对对应轮次、币种、投前/投后口径、融资金额及关键条款。"
            )
        elif kinds & {"revenue_actual", "revenue_projection"}:
            why = (
                "This scope directly affects revenue quality, growth analysis and valuation multiples."
                if is_english
                else "该口径会直接影响收入质量、增长判断与估值倍数。"
            )
            action = (
                "Reconcile the relevant period, revenue-recognition basis, order support and forecast assumptions."
                if is_english
                else "核对对应期间、收入确认口径、订单支撑和预测假设。"
            )
        elif "ip" in kinds:
            why = (
                "This scope affects control over core IP and the assessment of technical defensibility."
                if is_english
                else "该口径会影响核心知识产权的可控制性与技术壁垒判断。"
            )
            action = (
                "Verify ownership, legal status, licensing/transfer arrangements and the corresponding patent list."
                if is_english
                else "核对权属、法律状态、许可/转让安排及对应专利清单。"
            )
        else:
            why = (
                "This unresolved scope affects commercialization or capacity analysis."
                if is_english
                else "该未解决口径会影响商业化或产能判断。"
            )
            action = (
                "Verify related contracts, project status, entity relationships, timing basis and realization conditions."
                if is_english
                else "核对关联合同、项目状态、主体关系、时间口径与实现条件。"
            )

        version_terms = (
            ("round", "version", "new round", "previous", "earlier")
            if is_english
            else ("轮次", "版本", "新一轮", "此前")
        )
        flag_type = (
            "unresolved_version"
            if any(term.lower() in text.lower() for term in version_terms)
            else "scope_or_definition_gap"
        )

        flags.append({
            "title": (
                f"{title}: material scope still requires confirmation"
                if is_english
                else f"{title}仍有关键口径待确认"
            ),
            "type": flag_type,
            "unresolved_question": (
                "The supplied materials still do not resolve the following scope issue: "
                + text[:260]
                if is_english
                else "现有材料仍未解决以下口径问题：" + text[:260]
            ),
            "why": why,
            "action": action,
            "fact_ids": ids,
        })
        existing_ids.append(current_set)
        added += 1

    return output, added


def clean_research_output_refs(output, claims):
    valid_ids = {
        item["claim_id"]
        for item in claims
    }
    claim_map = {
        item["claim_id"]: item
        for item in claims
    }

    def clean_ids(values):
        return [
            value for value in (values or [])
            if value in valid_ids
        ]

    def ids_are_summary_eligible(ids):
        items = [
            claim_map[fact_id]
            for fact_id in ids
            if fact_id in claim_map
        ]

        if not items:
            return False

        return all(
            item.get("evidence_check", {}).get("status") == "pass"
            and item.get("evidence_status") != "Subjective / Marketing"
            for item in items
        )

    def numbers_are_supported(text, ids):
        requested = set(numeric_tokens(text))
        if not requested:
            return True
        evidence = " ".join(
            claim_map[fact_id].get("source_text", "")
            + " "
            + claim_map[fact_id].get("claim", "")
            for fact_id in ids
            if fact_id in claim_map
        )
        available = set(numeric_tokens(evidence))
        return requested.issubset(available)

    cleaned = {
        "headline": "",
        "executive_summary": [],
        "key_metrics": [],
        "sections": [],
        "research_flags": [],
    }

    for item in output.get("executive_summary", []):
        ids = clean_ids(item.get("fact_ids"))
        text = str(item.get("text", "")).strip()
        if (
            text
            and ids
            and ids_are_summary_eligible(ids)
            and numbers_are_supported(text, ids)
        ):
            cleaned["executive_summary"].append({
                "text": text,
                "fact_ids": ids,
            })

    for item in output.get("key_metrics", []):
        ids = clean_ids(item.get("fact_ids"))
        label = str(item.get("label", "")).strip()
        value = str(item.get("value", "")).strip()
        if (
            label
            and value
            and ids
            and numeric_tokens(value)
            and ids_are_summary_eligible(ids)
            and numbers_are_supported(value, ids)
        ):
            cleaned["key_metrics"].append({
                "label": label,
                "value": value,
                "note": str(item.get("note", "")).strip(),
                "fact_ids": ids,
            })

    for section in output.get("sections", []):
        title = str(section.get("title", "")).strip()
        if not title:
            continue

        groups = []
        for group in section.get("groups", []):
            ids = clean_ids(group.get("fact_ids"))
            group_title = str(group.get("title", "")).strip()
            group_summary = str(group.get("summary", "")).strip()
            if group_summary and not numbers_are_supported(group_summary, ids):
                group_summary = ""

            if group_title and ids:
                groups.append({
                    "title": group_title,
                    "summary": group_summary,
                    "fact_ids": ids,
                })

        if groups:
            cleaned["sections"].append({
                "title": title,
                "summary": str(section.get("summary", "")).strip(),
                "groups": groups,
            })

    for flag in output.get("research_flags", []):
        ids = clean_ids(flag.get("fact_ids"))
        title = str(flag.get("title", "")).strip()
        flag_type = str(flag.get("type", "")).strip()
        unresolved_question = str(
            flag.get("unresolved_question", "")
        ).strip()
        why = str(flag.get("why", "")).strip()
        action = str(flag.get("action", "")).strip()

        if (
            title
            and ids
            and flag_type in RESEARCH_FLAG_TYPES
            and unresolved_question
            and why
            and action
        ):
            cleaned["research_flags"].append({
                "title": title,
                "type": flag_type,
                "unresolved_question": unresolved_question,
                "why": why,
                "action": action,
                "fact_ids": ids,
            })

    cleaned["sections"] = filter_matrix_primary_home_sections(
        cleaned["sections"],
        claim_map,
    )

    cleaned["key_metrics"] = prune_key_metrics(
        cleaned["key_metrics"],
        claim_map,
    )

    metric_fact_sets = [
        set(item.get("fact_ids", []))
        for item in cleaned["key_metrics"]
    ]

    cleaned["executive_summary"] = [
        item
        for item in cleaned["executive_summary"]
        if not (
            len(item.get("fact_ids", [])) == 1
            and any(
                set(item["fact_ids"]).issubset(metric_ids)
                for metric_ids in metric_fact_sets
            )
        )
    ]

    return cleaned


def synthesize_research_output(
    client,
    claims,
    model_name,
    partial_scope=False,
    output_language="zh-CN",
):
    facts = [
        compact_fact_for_model(item)
        for item in claims
    ]

    scope_instruction = (
        "本次是快速抽样测试，只覆盖部分材料单元。禁止根据未抽样页面的缺失"
        "生成“全文未披露/没有提供/缺少独立验证”等 material_evidence_gap。"
        "只有当前已抽取 facts 自身形成的冲突、定义/口径问题、集合关系问题，"
        "或同一抽样范围内可直接证明的未解决问题，才可以进入 Flags。"
        if partial_scope
        else "本次按完整分析范围生成研究结果。"
    )

    language_instruction = (
        """
All user-visible generated fields must be written in clear professional English:
executive_summary.text; key_metrics label/value/note; group title/summary; and
Research Flag title/unresolved_question/why/action.

Keep JSON keys, fact_ids, evidence enums, Research Flag type enums and the
top-level section.title values exactly in the canonical internal schema.
section.title must remain one of the Chinese canonical domain values specified
below even when visible content is English; the UI translates those domain
labels separately.

Do not translate or rewrite source_text quotations. Preserve source-language
evidence exactly. Do not rescale numeric values during translation: preserve
the same numeric tokens used by supporting facts and translate unit wording
instead. Preserve dates in numeric form as well (for example, use 2021-05
rather than May 2021). This keeps deterministic numeric evidence checks stable.
"""
        if output_language == "en"
        else """
所有用户可见生成字段使用简体中文。JSON keys、fact_ids、证据状态枚举、
Research Flag type 枚举以及规定的顶层 section.title 内部值保持不变。
source_text 原文证据不得翻译或改写。
"""
    )

    prompt = f"""
你是一名投资研究整理助手。下面已经是经过来源约束的结构化 facts。

分析范围约束：{scope_instruction}

输出语言约束：
{language_instruction}

你的目标不是继续制造更多信息，而是把 facts 压缩成一个研究员 2 分钟可以读懂、需要时再深入的研究结果。

核心原则：
1. 不要逐条复述所有 facts。
2. 把相同主体、相同指标或同一业务事项合并。
3. 收入按年份合并；融资与估值合并；订单/意向按客户和状态合并；技术参数按产品/技术路线合并；产线进度按同一项目合并。
4. 不相关的信息绝对不要为了“主题相似”硬放到一起。
5. executive_summary 和 key_metrics 只能使用 evidence_check=pass 的 facts。
6. 预测必须保留“预计/计划/目标”，公司自述不得改写成独立验证事实。
7. 不要自行创造总额。只有当若干 fact 明确属于同一口径、同一币种且可以直接相加时，才可以写“已披露合计”，并必须引用全部组成 fact_ids；否则分别展示。
8. 每一句总结都必须带 fact_ids。没有 fact_ids 的内容不要写。
9. 技术参数不要逐个平铺，尽量合并为 1–3 个真正有研究价值的产品/技术组。
10. Key Metrics 只放真正改变投资判断的数字型指标，例如收入、融资、估值、订单金额、产能、专利数量和少量关键性能数值。每个 key_metric 只表达一个主要数字/一个口径，value 不要用分号把多个估值、多个技术参数或多个对象塞在同一张卡里；需要两个数字时拆成两项。同一 fact_ids / 同一技术主题默认最多贡献 1 个 Key Metric；技术类最多 2 项。公司自身商业/财务/产能数字优先于竞争者对标；Xanadu、PsiQuantum 等外部对标默认留在 Matrix，不进入核心数字，除非公司自身数字不足且对标本身就是当前核心研究问题。设备台数、型号编号、一般技术参数不要因为“有数字”就进入 Key Metrics；材料以技术为主时优先集成规模、吞吐/产能、迭代效率或良率。
11. Executive Summary 不要机械复述单个 Key Metric；每个 bullet 只回答一个投资问题，不要把融资、客户、产能、技术等互不相关的事实硬并成一句。优先顺序是财务/融资/客户订单 → 商业化/产能 → 团队/IP → 技术；只有材料确实主要是技术时才以技术为主。Subjective / Marketing 不应作为 Executive Summary 的核心结论。
12. Matrix 和 Flags 的职责严格分开：Matrix 保存“值得知道的现有信息”；Flags 只保存“尚未解决、会影响投资判断、且研究员有明确下一步动作”的问题。Flag 必须同时满足 Material + Unresolved + Actionable。单纯 Unverified、普通营销措辞、正常时间推进、以及已经明确解决的版本更新都不能成为 Flag。
13. 技术“可用于某行业”不等于已经形成该行业业务。application_scenario 必须使用“应用场景/方案/潜在用途”等措辞；只有存在 customer_reference、订单、合同、收入或真实部署证据时，才可以写“客户/业务/商业化”。
14. 所有用户可见文字里禁止出现 claim_1、fact_2 等内部 ID；fact_ids 只能放在结构化字段里。
15. 压缩同一主题时必须保留会改变投资解读的重要限定条件，例如“合作开发”“尚未授权”“采购意向而非已签约”“规划产能而非已实现产能”。如果这些限定来自同组其他 facts，summary 必须同时引用相关 fact_ids，不能只选一条看起来最完整的 claim。
16. 同一 comparison_key / period 出现明确的“调整、上调、下调、更新、修订”关系时，Research Matrix 的当前整理必须表达最新已知口径，并说明由此前什么口径变化而来；不得把已被明确替代的旧预测继续写成当前状态。
17. 如果材料不足以判断哪个版本更新，必须写成“存在多个版本/需核对”，并可生成 unresolved_version Flag；如果材料已经明确说明“调整/更新/修订”关系，则只更新 Matrix 当前状态，不要因为“发生过 revision”本身生成 Flag。
18. sections 只使用并按以下顺序组织：公司概览 → 商业化与竞争 → 财务与融资 → 团队与组织 → 产品与技术 → 知识产权 → 其他重要事项。商业化、客户、订单、项目、产能、市场定位和竞争格局都归入“商业化与竞争”，不要拆成互相分散的顶层领域。
19. 团队不是附属信息。“团队与组织”至少区分核心团队、人员规模/人员结构、研发/销售等关键职能配置；材料没有相关事实时才省略。
20. 知识产权不要只压成“有多少项专利”。如果 facts 支持，应拆成专利组合、授权状态、权属与合作开发、核心技术覆盖等不同 topic；但不要制造材料中没有的法律结论。
21. 同一可比事项随时间推进时，当前整理优先写最新已知状态，并在有研究价值时简短保留前一关键节点。例如3月27个项目、6月30个项目，应写当前30个并说明较3月增加，而不是随机保留旧值。
22. 不要为了让界面看起来短而删除高价值 topic，但每组事实要有一个 primary home：同一组 facts 不要为了“完整”同时复制到产品与技术、知识产权、其他重要事项等多个顶层领域。Summary 和 Flags 可以引用 Matrix 事实，但 Matrix 本身避免跨领域重复。
23. 禁止把可能重叠的分类当成互斥集合做“剩余量”推导。除非材料明确说明分类互斥且穷尽，否则不能用“总数 - 子类A - 子类B”推导第三类数量。例如11项专利、5项合作开发、3项申请/开发中，不能推出3项为公司独立持有且已授权。
24. Research Flag 的 type 只能是 unresolved_conflict / unresolved_version / scope_or_definition_gap / material_evidence_gap。每个 Flag 必须填写 unresolved_question，明确写出“现有材料还不能回答什么”，同时填写 why 和 action。若没有清楚的 unresolved question 或没有明确可执行的下一步，就不要生成 Flag。
25. 可重点检查但不能机械报警的 Flag 领域包括：客户/订单真实性与状态、客户或供应商集中度及关键依赖、收入质量与预测假设、融资/估值口径与关键条款、产能/量产/交付真实性、知识产权权属/授权/许可、核心团队或关键人依赖、治理/关联交易、监管资质/合规/诉讼/政策依赖、市场份额/市场规模/“唯一第一”等竞争主张的依据。
26. “材料没有写”本身不是 Flag。只有当现有 facts 已经表明某个问题对投资判断重要，或者某个核心主张依赖缺失信息才能成立时，缺口才可成为 material_evidence_gap。不要把通用尽调清单里的每一项都因为未披露而自动报警。
27. 如果分析范围约束说明这是快速抽样测试，不得把“当前抽样 facts 没看到某信息”升级为整份材料的缺失结论。融资金额、估值、客户合同、团队人数、第三方验证等“可能存在于未抽样页面”的缺失项，在 smoke test 中默认不生成 material_evidence_gap；结构性歧义（例如分类是否重叠）、明确冲突和强竞争主张的定义口径仍可生成 Flag。
28. 如果 Executive Summary 或 Matrix 已经明确写出“需核对 / 口径不明 / 待确认 / 无法确认”等会影响投资判断的实质 caveat，必须同步生成对应 Research Flag；不要让系统在摘要里发现了问题，却没有进入待办。

严格返回 JSON 对象：
{{
  "headline": "",
  "executive_summary": [
    {{"text": "一句高密度总结", "fact_ids": ["fact_1"]}}
  ],
  "key_metrics": [
    {{
      "label": "2025收入目标",
      "value": "3亿元",
      "note": "预测/目标",
      "fact_ids": ["fact_2"]
    }}
  ],
  "sections": [
    {{
      "title": "财务与融资",
      "summary": "本节一行总结",
      "groups": [
        {{
          "title": "收入规划",
          "summary": "把同类事实合并后的短总结",
          "fact_ids": ["fact_2", "fact_3"]
        }}
      ]
    }}
  ],
  "research_flags": [
    {{
      "title": "同一合同金额存在多个未解决口径",
      "type": "unresolved_conflict|unresolved_version|scope_or_definition_gap|material_evidence_gap",
      "unresolved_question": "现有材料还不能回答的具体问题",
      "why": "为什么这个未解决问题会影响投资判断",
      "action": "研究员下一步应该核查什么",
      "fact_ids": ["fact_4", "fact_5"]
    }}
  ]
}}

控制规模：
- executive_summary：最多 6 条
- key_metrics：最多 10 项
- sections：最多 7 个，按规定的自上而下顺序输出
- 每个 section groups：最多 8 个；不要因为界面折叠而主动丢掉高价值主题
- research_flags：最多 6 个

facts：
{json.dumps(facts, ensure_ascii=False)}
"""

    record_api_call("synthesis")
    response = client.responses.create(
        model=model_name,
        input=prompt,
    )
    result = parse_model_json(response.output_text)

    if not isinstance(result, dict):
        raise ValueError("研究汇总必须返回 JSON 对象。")

    return clean_research_output_refs(
        result,
        claims,
    )


def verification_risk_score(
    target_kind,
    text,
    fact_ids,
    claim_map,
):
    items = [
        claim_map[fact_id]
        for fact_id in fact_ids
        if fact_id in claim_map
    ]
    if not items:
        return -1

    score = 0

    if target_kind == "flag":
        return 100
    if target_kind == "summary":
        score += 6
    if len(items) > 1:
        score += 5
    if any(
        item.get("evidence_status")
        in {
            "Projection",
            "Subjective / Marketing",
        }
        for item in items
    ):
        score += 5
    if any(
        item.get(
            "evidence_check",
            {},
        ).get("status") == "review"
        for item in items
    ):
        score += 3
    if numeric_tokens(text):
        score += 2
    if any(
        keyword in str(text)
        for keyword in (
            "调整",
            "更新",
            "修订",
            "其中",
            "合作开发",
            "未授权",
            "唯一",
            "第一",
            "领先",
            "revised",
            "updated",
            "co-developed",
            "not granted",
            "only",
            "leading",
        )
    ):
        score += 3
    if target_kind == "metric":
        score -= 2

    return score


def build_verification_items(
    research_output,
    claims,
    limit=14,
):
    claim_map = {
        item["claim_id"]: item
        for item in claims
    }
    candidates = []

    def add_candidate(
        target,
        target_kind,
        text,
        fact_ids,
        order,
    ):
        fact_ids = [
            fact_id
            for fact_id in (fact_ids or [])
            if fact_id in claim_map
        ]
        if not fact_ids:
            return

        score = verification_risk_score(
            target_kind,
            text,
            fact_ids,
            claim_map,
        )

        if (
            target_kind in {"group", "metric"}
            and score < 5
        ):
            return

        candidates.append({
            "target": target,
            "text": text,
            "fact_ids": fact_ids,
            "_risk": score,
            "_order": order,
        })

    order = 0

    for index, item in enumerate(
        research_output.get(
            "research_flags",
            [],
        )
    ):
        add_candidate(
            f"flag:{index}",
            "flag",
            " | ".join(
                part
                for part in [
                    item.get("title", ""),
                    item.get(
                        "unresolved_question",
                        "",
                    ),
                    item.get("why", ""),
                    item.get("action", ""),
                ]
                if part
            ),
            item.get("fact_ids", []),
            order,
        )
        order += 1

    for index, item in enumerate(
        research_output.get(
            "executive_summary",
            [],
        )
    ):
        add_candidate(
            f"summary:{index}",
            "summary",
            item.get("text", ""),
            item.get("fact_ids", []),
            order,
        )
        order += 1

    for section_index, section in enumerate(
        research_output.get(
            "sections",
            [],
        )
    ):
        for group_index, group in enumerate(
            section.get("groups", [])
        ):
            text = str(
                group.get("summary", "")
            ).strip()
            if not text:
                continue
            add_candidate(
                (
                    f"group:{section_index}:"
                    f"{group_index}"
                ),
                "group",
                text,
                group.get("fact_ids", []),
                order,
            )
            order += 1

    for index, item in enumerate(
        research_output.get(
            "key_metrics",
            [],
        )
    ):
        add_candidate(
            f"metric:{index}",
            "metric",
            (
                f"{item.get('label', '')}: "
                f"{item.get('value', '')} "
                f"{item.get('note', '')}"
            ).strip(),
            item.get("fact_ids", []),
            order,
        )
        order += 1

    candidates.sort(
        key=lambda item: (
            -item["_risk"],
            item["_order"],
        )
    )

    return [
        {
            key: value
            for key, value in item.items()
            if not key.startswith("_")
        }
        for item in candidates[:limit]
    ]

def apply_verification_issues(
    research_output,
    issues,
):
    audited = json.loads(
        json.dumps(
            research_output,
            ensure_ascii=False,
        )
    )

    drop_summary = set()
    drop_metrics = set()
    drop_groups = defaultdict(set)
    drop_flags = set()

    for issue in issues:
        if not isinstance(issue, dict):
            continue

        target = str(
            issue.get("target", "")
        ).strip()
        decision = str(
            issue.get("decision", "")
        ).strip()
        corrected_text = str(
            issue.get("corrected_text", "")
        ).strip()

        parts = target.split(":")
        if not parts:
            continue

        try:
            if parts[0] == "summary" and len(parts) == 2:
                index = int(parts[1])
                if decision == "drop":
                    drop_summary.add(index)
                elif decision == "replace" and corrected_text:
                    audited["executive_summary"][index][
                        "text"
                    ] = corrected_text

            elif parts[0] == "metric" and len(parts) == 2:
                index = int(parts[1])
                if decision == "drop":
                    drop_metrics.add(index)

            elif parts[0] == "group" and len(parts) == 3:
                section_index = int(parts[1])
                group_index = int(parts[2])

                if decision == "drop":
                    drop_groups[section_index].add(
                        group_index
                    )
                elif decision == "replace" and corrected_text:
                    audited["sections"][section_index][
                        "groups"
                    ][group_index]["summary"] = (
                        corrected_text
                    )

            elif parts[0] == "flag" and len(parts) == 2:
                index = int(parts[1])
                if decision == "drop":
                    drop_flags.add(index)

        except (
            ValueError,
            IndexError,
            KeyError,
            TypeError,
        ):
            continue

    audited["executive_summary"] = [
        item
        for index, item in enumerate(
            audited.get("executive_summary", [])
        )
        if index not in drop_summary
    ]

    audited["key_metrics"] = [
        item
        for index, item in enumerate(
            audited.get("key_metrics", [])
        )
        if index not in drop_metrics
    ]

    for section_index, indexes in drop_groups.items():
        if section_index >= len(
            audited.get("sections", [])
        ):
            continue

        groups = audited["sections"][
            section_index
        ].get("groups", [])

        audited["sections"][
            section_index
        ]["groups"] = [
            item
            for index, item in enumerate(groups)
            if index not in indexes
        ]

    audited["sections"] = [
        section
        for section in audited.get("sections", [])
        if section.get("groups")
    ]

    audited["research_flags"] = [
        item
        for index, item in enumerate(
            audited.get("research_flags", [])
        )
        if index not in drop_flags
    ]

    return audited


def verify_research_output(
    client,
    research_output,
    claims,
    model_name,
    output_language="zh-CN",
):
    audit_items = build_verification_items(
        research_output,
        claims,
    )

    if not audit_items:
        return clean_research_output_refs(
            research_output,
            claims,
        )
    referenced_ids = {
        fact_id
        for item in audit_items
        for fact_id in item.get("fact_ids", [])
    }

    fact_map = {
        item["claim_id"]: compact_fact_for_model(item)
        for item in claims
        if item["claim_id"] in referenced_ids
    }

    verification_language_instruction = (
        "When decision=replace, corrected_text must remain in professional English."
        if output_language == "en"
        else "decision=replace 时，corrected_text 保持简体中文。"
    )

    prompt = f"""
你是 evidence verifier。只找问题，不要重写整份研究结果。

输出语言约束：{verification_language_instruction}

audit_items 已经由 Python 按风险筛过，只检查这些高风险输出；不要要求补齐未列入 audit_items 的普通条目。
请逐项检查 audit_items 是否被其 fact_ids 和 facts 直接支持。

只检查这些问题：
- 是否把“公司称/材料称”增强成客观事实；
- 是否把 Projection 写成已实现；
- 是否改变数字、日期、主体或范围；
- 是否写入 fact_ids 无法支持的新信息；
- 是否把营销/主观表述伪装成独立验证事实；
- 是否遗漏同一组 fact_ids 中会改变解读的重要限定条件，例如合作开发、尚未授权、采购意向与已签约的区别；
- 是否把已经被明确调整/更新/修订的旧版本继续写成当前状态；
- 是否在同一可比事项已有更晚时间点时，仍把较早值写成“当前状态”；
- 是否把已经明确解决的时间推进或版本调整错误保留成 Research Flag；
- 是否通过可能重叠的分类做了无证据的剩余量推导，例如“11总量 - 5合作开发 - 3申请中 = 3独立持有且已授权”；
- 每个 Flag 是否真的包含现有材料无法回答的 unresolved_question，并且这个问题会影响投资判断且存在明确下一步动作。

如果一项完全成立，不要输出它。
如果只有措辞需要收紧，decision=replace，并只给 corrected_text。
如果整项缺乏支持，decision=drop。
metric 只允许 keep 或 drop；有问题时直接 drop。
flag 本身可以表达“待核查”，但不能断言未被支持的事实；有问题时直接 drop。

严格只返回：
{{
  "issues": [
    {{
      "target": "summary:0 或 metric:1 或 group:0:2 或 flag:1",
      "decision": "replace|drop",
      "corrected_text": "仅 replace 时填写"
    }}
  ]
}}

不要返回原始 research_output，不要解释，不要新增 target。

audit_items：
{json.dumps(audit_items, ensure_ascii=False)}

facts：
{json.dumps(fact_map, ensure_ascii=False)}
"""

    record_api_call("verify")
    response = client.responses.create(
        model=model_name,
        input=prompt,
    )
    result = parse_model_json(response.output_text)

    if not isinstance(result, dict):
        raise ValueError("验证器必须返回 JSON 对象。")

    issues = result.get("issues", [])
    if not isinstance(issues, list):
        raise ValueError("验证器 issues 必须是 JSON 数组。")

    audited = apply_verification_issues(
        research_output,
        issues,
    )

    return clean_research_output_refs(
        audited,
        claims,
    )

def parse_cny_amount(text):
    matches = re.findall(
        r"(\d+(?:\.\d+)?)\s*(亿元|万元|元)",
        str(text or ""),
    )
    if len(matches) != 1:
        return None

    value, unit = matches[0]
    value = float(value)

    multiplier = {
        "亿元": 100_000_000,
        "万元": 10_000,
        "元": 1,
    }[unit]

    return value * multiplier


def format_cny_amount(amount, output_language="zh-CN"):
    if amount >= 100_000_000:
        value = amount / 100_000_000
        return (
            f"RMB {value:g} hundred million"
            if output_language == "en"
            else f"{value:g}亿元"
        )
    if amount >= 10_000:
        value = amount / 10_000
        return (
            f"RMB {value:g} ten-thousand"
            if output_language == "en"
            else f"{value:g}万元"
        )
    return (
        f"RMB {amount:g}"
        if output_language == "en"
        else f"{amount:g}元"
    )


def compute_disclosed_totals(claims, output_language="zh-CN"):
    configs = [
        (
            "order_signed",
            (
                "Directly summable disclosed signed-order value"
                if output_language == "en"
                else "可直接合计的已披露签约金额"
            ),
            (
                "Includes only evidence-pass signed orders without duplicate comparison keys; ambiguous versions/scopes are excluded."
                if output_language == "en"
                else "只合计证据检查通过且未出现重复 comparison_key 的已签约订单；存在版本/口径疑问的合同不会被机械相加。"
            ),
        ),
        (
            "purchase_intent",
            (
                "Directly summable purchase-intent value"
                if output_language == "en"
                else "可直接合计的采购意向金额"
            ),
            (
                "Includes purchase intents only; signed contracts are not mixed in, and ambiguous versions/scopes are excluded."
                if output_language == "en"
                else "仅合计采购意向，不与已签约合同混合；存在版本/口径疑问的项目不进入总额。"
            ),
        ),
    ]

    metrics = []

    for fact_kind, label, note in configs:
        candidates = []
        comparison_counts = defaultdict(int)

        for item in claims:
            if item.get("fact_kind") != fact_kind:
                continue
            if item.get("evidence_check", {}).get("status") != "pass":
                continue

            amount = parse_cny_amount(item.get("claim", ""))
            if amount is None:
                amount = parse_cny_amount(item.get("source_text", ""))
            if amount is None:
                continue

            comparison_key = item.get("comparison_key", "").strip()
            if comparison_key:
                comparison_counts[comparison_key] += 1

            candidates.append((item, amount))

        # 同一 comparison_key 出现多条时，可能是版本、口径或重复材料。
        # 为避免把潜在重复合同机械相加，自动总额中直接排除这些组。
        unique = {}
        for item, amount in candidates:
            comparison_key = item.get("comparison_key", "").strip()
            if comparison_key and comparison_counts[comparison_key] > 1:
                continue

            dedupe_key = (
                item.get("entity", ""),
                item.get("period", ""),
                item.get("topic_group", ""),
                amount,
            )
            unique.setdefault(
                dedupe_key,
                (amount, item["claim_id"]),
            )

        if len(unique) < 2:
            continue

        total = sum(value for value, _ in unique.values())
        fact_ids = [
            fact_id
            for _, fact_id in unique.values()
        ]

        metrics.append({
            "label": label,
            "value": format_cny_amount(
                total,
                output_language=output_language,
            ),
            "note": note,
            "fact_ids": fact_ids,
        })

    return metrics


def merge_computed_metrics(
    research_output,
    claims,
    output_language="zh-CN",
):
    metrics = list(research_output.get("key_metrics", []))
    existing_labels = {
        item.get("label")
        for item in metrics
    }

    for metric in compute_disclosed_totals(
        claims,
        output_language=output_language,
    ):
        if metric["label"] not in existing_labels:
            metrics.insert(0, metric)

    research_output["key_metrics"] = metrics[:10]
    return research_output


def clean_user_text(value):
    text = str(value or "").strip()
    text = re.sub(
        r"(?:claim|fact)_\d+\s*称?",
        "",
        text,
        flags=re.I,
    )
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip(" ：，,;")


def display_text_key(value):
    text = clean_user_text(value)
    text = re.sub(r"\s+", "", text)
    return text.strip("。；;，,：:")


def write_source_text_if_distinct(item):
    source_text = str(item.get("source_text", "") or "").strip()
    if (
        source_text
        and display_text_key(source_text)
        != display_text_key(item.get("claim", ""))
    ):
        st.write(source_text)


def analyst_review_key(kind, title, items):
    fingerprints = sorted(
        "|".join(
            [
                str(item.get("source_filename", "")),
                str(item.get("source_location", "")),
                str(item.get("claim", "")),
            ]
        )
        for item in (items or [])
    )
    payload = (
        str(kind)
        + "|"
        + str(title)
        + "|"
        + "||".join(fingerprints)
    )
    return hashlib.sha256(
        payload.encode("utf-8")
    ).hexdigest()[:20]


def matrix_review_key(domain, label, items):
    return analyst_review_key(
        "matrix",
        f"{domain}|{label}",
        items,
    )


def flag_review_key(flag, related_items):
    return analyst_review_key(
        "flag",
        (
            str(flag.get("title", ""))
            + "|"
            + str(flag.get("type", ""))
        ),
        related_items,
    )


def analyst_review_status_label(status, language=None):
    return review_status_label(
        status,
        language or current_ui_language(),
    )


def analyst_review_status_from_label(label):
    if label in ANALYST_REVIEW_OPTIONS:
        return label

    for status in ANALYST_REVIEW_OPTIONS:
        if label in {
            review_status_label(status, "zh-CN"),
            review_status_label(status, "en"),
        }:
            return status
    return "unreviewed"


def get_analyst_review(review_key):
    reviews = st.session_state.setdefault(
        "analyst_reviews",
        {},
    )
    stored = reviews.get(
        review_key,
        {
            "status": "unreviewed",
            "note": "",
        },
    )

    status_widget_key = (
        f"analyst_review_status_{review_key}"
    )
    note_widget_key = (
        f"analyst_review_note_{review_key}"
    )

    status = stored.get(
        "status",
        "unreviewed",
    )
    if status_widget_key in st.session_state:
        status = analyst_review_status_from_label(
            st.session_state[status_widget_key]
        )

    note = stored.get("note", "")
    if note_widget_key in st.session_state:
        note = st.session_state[note_widget_key]

    resolved = {
        "status": status,
        "note": str(note or "").strip(),
    }
    reviews[review_key] = resolved
    return resolved


def render_analyst_review_controls(review_key):
    reviews = st.session_state.setdefault(
        "analyst_reviews",
        {},
    )
    current = get_analyst_review(review_key)

    status_widget_key = (
        f"analyst_review_status_{review_key}"
    )
    note_widget_key = (
        f"analyst_review_note_{review_key}"
    )

    if status_widget_key not in st.session_state:
        st.session_state[status_widget_key] = (
            analyst_review_status_label(
                current["status"]
            )
        )
    if note_widget_key not in st.session_state:
        st.session_state[note_widget_key] = (
            current.get("note", "")
        )

    if (
        st.session_state.get(status_widget_key)
        not in ANALYST_REVIEW_OPTIONS
    ):
        st.session_state[status_widget_key] = current["status"]

    status = st.radio(
        ui("人工复核", "Analyst review"),
        ANALYST_REVIEW_OPTIONS,
        format_func=lambda value: analyst_review_status_label(value),
        horizontal=True,
        key=status_widget_key,
    )
    note = st.text_input(
        ui("复核备注（可选）", "Review note (optional)"),
        key=note_widget_key,
        placeholder=ui(
            "例如：核对合作开发专利的权属；或记录为什么暂不需要跟进。",
            "For example: verify ownership of co-developed patents, or record why no follow-up is currently needed.",
        ),
    )
    reviews[review_key] = {
        "status": status,
        "note": str(note or "").strip(),
    }

    if status == "cleared":
        st.success(
            ui(
                "已记录：研究员已看过，当前暂不需要进一步跟进。",
                "Recorded: reviewed by the analyst; no further follow-up is currently required.",
            )
        )
    elif status == "follow_up":
        st.warning(
            ui(
                "已记录：该主题仍需要进一步尽调或确认。",
                "Recorded: this topic still requires further diligence or confirmation.",
            )
        )

    st.caption(
        ui(
            "人工复核只记录研究员的处理状态，不会把 Unverified 自动改成 Supported。",
            "Analyst review records workflow state only; it never converts Unverified evidence into Supported evidence.",
        )
    )


def export_analyst_reviews():
    reviews = st.session_state.get(
        "analyst_reviews",
        {},
    )
    return {
        key: {
            "status": value.get(
                "status",
                "unreviewed",
            ),
            "status_label": analyst_review_status_label(
                value.get(
                    "status",
                    "unreviewed",
                )
            ),
            "note": value.get("note", ""),
        }
        for key, value in reviews.items()
        if (
            value.get("status", "unreviewed")
            != "unreviewed"
            or value.get("note")
        )
    }


def research_flag_entries(
    research_output,
    claim_map,
):
    entries = []

    for flag in research_output.get(
        "research_flags",
        [],
    ):
        if not flag_is_actionable(
            flag,
            claim_map,
        ):
            continue

        related = [
            claim_map[fact_id]
            for fact_id in flag.get(
                "fact_ids",
                [],
            )
            if fact_id in claim_map
        ]
        review_key = flag_review_key(
            flag,
            related,
        )
        review = get_analyst_review(
            review_key
        )
        entries.append({
            "flag": flag,
            "related": related,
            "review_key": review_key,
            "review": review,
        })

    return entries


def compact_location(location):
    location = str(location or "").strip()
    page = re.fullmatch(r"Page\s+(\d+)", location, flags=re.I)
    if page:
        return f"P{page.group(1)}"

    slide = re.fullmatch(r"Slide\s+(\d+)", location, flags=re.I)
    if slide:
        return f"S{slide.group(1)}"

    paragraphs = re.fullmatch(
        r"Paragraphs\s+(\d+)-(\d+)",
        location,
        flags=re.I,
    )
    if paragraphs:
        return f"¶{paragraphs.group(1)}-{paragraphs.group(2)}"

    lines = re.fullmatch(
        r"Lines\s+(\d+)-(\d+)",
        location,
        flags=re.I,
    )
    if lines:
        return f"L{lines.group(1)}-{lines.group(2)}"

    return location


def snapshot_material_profile(claims, sections):
    section_names = [
        domain_label(
            section.get("title", ""),
            current_ui_language(),
        )
        for section in sections
        if section.get("title")
    ][:3]

    marketing = sum(
        item.get("evidence_status") == "Subjective / Marketing"
        for item in claims
    )
    needs_review = sum(
        item.get("evidence_check", {}).get("status") == "review"
        for item in claims
    )
    projection = sum(
        item.get("evidence_status") == "Projection"
        for item in claims
    )

    parts = []

    if section_names:
        parts.append(
            (
                "当前分析主要覆盖"
                + "、".join(section_names)
                + "。"
            )
            if current_ui_language() == "zh-CN"
            else
            "Current analysis primarily covers "
            + ", ".join(section_names)
            + "."
        )

    if marketing or needs_review or projection:
        detail = []
        if marketing:
            detail.append(
                f"{marketing}条公司/营销表述"
                if current_ui_language() == "zh-CN"
                else f"{marketing} company/marketing claims"
            )
        if projection:
            detail.append(
                f"{projection}条预测/目标"
                if current_ui_language() == "zh-CN"
                else f"{projection} projections/targets"
            )
        if needs_review:
            detail.append(
                f"{needs_review}条证据锚点待检查"
                if current_ui_language() == "zh-CN"
                else f"{needs_review} evidence anchors needing review"
            )
        parts.append(
            (
                "证据结构上，当前包含"
                + "、".join(detail)
                + "；证据锚点状态只表示程序匹配情况，不等于 Research Flag。"
            )
            if current_ui_language() == "zh-CN"
            else
            "Evidence structure includes "
            + ", ".join(detail)
            + "; evidence-anchor status reflects programmatic matching only and is not itself a Research Flag."
        )
    elif claims:
        parts.append(
            ui(
                "当前提取到的高价值事实均已保留来源，可在需要时回到证据核对。",
                "Extracted high-value facts retain source pointers and can be checked against evidence when needed.",
            )
        )

    return "".join(parts)


def research_flag_type_label(flag_type):
    return flag_type_label(
        str(flag_type or "").strip(),
        current_ui_language(),
    )


def flag_is_actionable(flag, claim_map):
    ids = [
        fact_id
        for fact_id in flag.get("fact_ids", [])
        if fact_id in claim_map
    ]
    if not ids:
        return False

    flag_type = str(flag.get("type", "")).strip()
    unresolved_question = str(
        flag.get("unresolved_question", "")
    ).strip()
    why = str(flag.get("why", "")).strip()
    action = str(flag.get("action", "")).strip()

    if flag_type not in RESEARCH_FLAG_TYPES:
        return False

    if not (
        unresolved_question
        and why
        and action
    ):
        return False

    # A Flag is a research task, not a synonym for "Unverified".
    # Resolved timelines/revisions belong in Matrix only.
    resolved_change_terms = (
        "已明确调整",
        "已更新为",
        "已修订为",
        "属于正常时间推进",
        "explicitly revised",
        "updated to",
        "revised to",
        "normal timeline progression",
    )
    flag_text = " ".join(
        [
            str(flag.get("title", "")),
            unresolved_question,
            why,
        ]
    )
    if any(
        term in flag_text
        for term in resolved_change_terms
    ):
        return False

    return True

def fact_sources(fact_ids, claim_map):
    seen = []
    for fact_id in fact_ids:
        item = claim_map.get(fact_id)
        if not item:
            continue
        source = (
            f"{item['source_filename']} · "
            f"{item['source_location']}"
        )
        if source not in seen:
            seen.append(source)
    return seen


def evidence_anchor_warning(items):
    return any(
        item.get("evidence_check", {}).get("status") == "review"
        for item in (items or [])
    )


def research_status(items):
    if not items:
        return "—"

    if any(
        item.get("evidence_status") == "Subjective / Marketing"
        for item in items
    ):
        return ui("◐ 公司/营销表述", "◐ Company / marketing claim")

    if any(
        item.get("evidence_status") == "Projection"
        for item in items
    ):
        return ui("◌ 预测 / 目标", "◌ Projection / target")

    if all(
        item.get("evidence_status") == "Supported"
        for item in items
    ):
        return ui("● 当前材料支持", "● Supported by supplied material")

    return ui("◐ 未独立验证", "◐ Not independently verified")


def compact_source_label(items, limit=3):
    filenames = {
        item.get("source_filename", "")
        for item in items
        if item.get("source_filename")
    }
    show_filename = len(filenames) > 1
    labels = []

    for item in items:
        filename = Path(item.get("source_filename", "")).stem
        location = compact_location(item.get("source_location", ""))

        if show_filename and filename:
            short_name = filename[:18] + ("…" if len(filename) > 18 else "")
            label = f"{short_name} · {location}"
        else:
            label = location

        if label and label not in labels:
            labels.append(label)

    if len(labels) <= limit:
        return " / ".join(labels)

    return " / ".join(labels[:limit]) + f" / +{len(labels) - limit}"


def investment_priority_from_items(items):
    kinds = {
        item.get("fact_kind", "other")
        for item in items
    }
    categories = {
        item.get("category", "其他")
        for item in items
    }

    if kinds & {
        "revenue_actual",
        "revenue_projection",
        "financing",
        "valuation",
    } or "财务与预测" in categories:
        return 0

    if kinds & {
        "order_signed",
        "purchase_intent",
        "customer_reference",
    } or "客户与订单" in categories:
        return 1

    if kinds & {
        "capacity",
        "production_plan",
    } or "商业化与产能" in categories:
        return 2

    if "ip" in kinds or "知识产权" in categories:
        return 3

    if "团队" in categories:
        return 4

    if kinds & {
        "technology_metric",
        "product",
    } or "技术与产品" in categories:
        return 5

    if "application_scenario" in kinds:
        return 6

    if "市场与竞争" in categories:
        return 7

    return 8


def summary_item_priority(item, claim_map):
    items = [
        claim_map[fact_id]
        for fact_id in item.get("fact_ids", [])
        if fact_id in claim_map
    ]
    return investment_priority_from_items(items)


def matrix_domain_priority(section_title, items):
    domain = canonical_matrix_domain(
        section_title,
        items,
    )
    try:
        return MATRIX_DOMAIN_ORDER.index(domain)
    except ValueError:
        return len(MATRIX_DOMAIN_ORDER)

def build_matrix_rows(research_output, claim_map):
    rows = []
    details = []

    for section in research_output.get("sections", []):
        for group in section.get("groups", []):
            items = [
                claim_map[fact_id]
                for fact_id in group.get("fact_ids", [])
                if fact_id in claim_map
            ]

            domain = canonical_matrix_domain(
                section.get("title", ""),
                items,
            )
            label = canonical_matrix_topic(
                group.get("title", ""),
                items,
                domain,
            )
            review_key = matrix_review_key(
                domain,
                label,
                items,
            )
            review = get_analyst_review(
                review_key
            )

            summary = (
                str(group.get("summary", "")).strip()
                or (items[0]["claim"] if items else "")
            )

            rows.append({
                "领域": domain_label(domain, current_ui_language()),
                "主题": label,
                "当前整理": clean_user_text(summary),
                "状态": research_status(items),
                "人工复核": analyst_review_status_label(
                    review["status"]
                ),
                "来源": (
                    compact_source_label(items)
                    + (" ⚠" if evidence_anchor_warning(items) else "")
                ),
            })

            details.append({
                "label": label,
                "section": domain,
                "group": group,
                "items": items,
                "review_key": review_key,
            })

    paired = list(zip(rows, details))
    paired.sort(
        key=lambda pair: (
            matrix_domain_priority(
                pair[1]["section"],
                pair[1]["items"],
            ),
            matrix_topic_priority(
                pair[1]["label"],
                pair[1]["items"],
                pair[1]["section"],
            ),
            pair[1]["label"],
        )
    )

    return (
        [pair[0] for pair in paired],
        [pair[1] for pair in paired],
    )

def source_sort_key(item):
    location = item.get("source_location", "")
    match = re.search(r"(\d+)", location)

    return (
        item.get("source_filename", ""),
        int(match.group(1)) if match else 10**9,
        location,
    )


def metric_group(metric, claim_map):
    items = [
        claim_map[fact_id]
        for fact_id in metric.get("fact_ids", [])
        if fact_id in claim_map
    ]
    kinds = {
        item.get("fact_kind", "other")
        for item in items
    }

    if kinds & {
        "revenue_actual",
        "revenue_projection",
        "financing",
        "valuation",
    }:
        return ui("商业与财务", "Commercial & financial"), 0

    if kinds & {
        "order_signed",
        "purchase_intent",
    }:
        return ui("客户与订单", "Customers & orders"), 1

    if kinds & {
        "capacity",
        "production_plan",
    }:
        return ui("产能与进展", "Capacity & progress"), 2

    if "ip" in kinds:
        return ui("知识产权", "Intellectual property"), 3

    if kinds & {
        "technology_metric",
        "product",
    }:
        return ui("技术指标", "Technical metrics"), 4

    return ui("其他指标", "Other metrics"), 5


def render_snapshot_tab(research_output, claims, claim_map):
    sections = research_output.get("sections", [])
    flag_entries = research_flag_entries(
        research_output,
        claim_map,
    )
    flags = [
        entry["flag"]
        for entry in flag_entries
        if entry["review"]["status"] != "cleared"
    ]
    summary = sorted(
        research_output.get("executive_summary", []),
        key=lambda item: summary_item_priority(
            item,
            claim_map,
        ),
    )
    metrics = research_output.get("key_metrics", [])

    locations = []
    for item in sorted(claims, key=source_sort_key):
        location = compact_location(
            item.get("source_location", "")
        )
        if location and location not in locations:
            locations.append(location)

    if locations:
        st.caption(
            ui("当前结果来源：", "Current result sources: ")
            + " / ".join(locations[:8])
            + (
                f" / +{len(locations) - 8}"
                if len(locations) > 8
                else ""
            )
        )

    if summary:
        st.markdown(ui("**What matters / 核心摘要**", "**What matters**"))
        for item in summary[:4]:
            text = clean_user_text(item.get("text", ""))
            if not text:
                continue
            st.write(f"• {text}")
            sources = fact_sources(
                item.get("fact_ids", []),
                claim_map,
            )
            if sources:
                st.caption(
                    ui("来源：", "Source: ")
                    + compact_source_label(
                        [
                            claim_map[fact_id]
                            for fact_id in item.get("fact_ids", [])
                            if fact_id in claim_map
                        ]
                    )
                )

    if not summary:
        profile = snapshot_material_profile(
            claims,
            sections,
        )
        if profile:
            with st.container(border=True):
                st.write(profile)

    if metrics:
        st.markdown(ui("**核心数字**", "**Core numbers**"))

        grouped_metrics = defaultdict(list)
        group_order = {}

        for metric in metrics:
            group_name, priority = metric_group(
                metric,
                claim_map,
            )
            grouped_metrics[group_name].append(metric)
            group_order[group_name] = priority

        shown = 0
        for group_name in sorted(
            grouped_metrics,
            key=lambda name: group_order[name],
        ):
            if shown >= 6:
                break

            group_items = grouped_metrics[group_name]
            remaining = 6 - shown
            group_items = group_items[:remaining]

            st.caption(group_name)
            for start in range(0, len(group_items), 3):
                cols = st.columns(3)
                for col, item in zip(
                    cols,
                    group_items[start:start + 3],
                ):
                    with col:
                        with st.container(border=True):
                            st.caption(
                                clean_user_text(item["label"])
                            )
                            st.markdown(
                                "**"
                                + clean_user_text(item["value"])
                                + "**"
                            )
                            if item.get("note"):
                                st.caption(
                                    clean_user_text(item["note"])
                                )
                        shown += 1

    if flags:
        st.markdown(ui("**需要注意**", "**Needs attention**"))
        for flag in flags[:3]:
            title = clean_user_text(flag.get("title", ""))
            flag_type = clean_user_text(
                research_flag_type_label(
                    flag.get("type", "")
                )
            )
            st.write(
                f"⚠ **{title}**"
                + (f" · {flag_type}" if flag_type else "")
            )
    else:
        st.caption(ui("当前没有识别到需要优先人工核查的问题。", "No issue currently requires priority manual review."))

def render_matrix_tab(research_output, claim_map):
    rows, details = build_matrix_rows(
        research_output,
        claim_map,
    )

    if not rows:
        st.caption(ui("当前没有可展示的研究主题。", "No research topics are available to display."))
        return

    section_options = ["全部"] + list(dict.fromkeys(
        item["section"]
        for item in details
    ))
    selected_section = st.selectbox(
        ui("领域", "Domain"),
        section_options,
        format_func=lambda value: (
            ui("全部", "All")
            if value == "全部"
            else domain_label(value, current_ui_language())
        ),
        key="research_matrix_section",
    )

    all_filtered_pairs = [
        (row, detail)
        for row, detail in zip(rows, details)
        if (
            selected_section == "全部"
            or detail["section"] == selected_section
        )
    ]

    display_pairs = all_filtered_pairs
    hidden_count = 0

    if selected_section == "全部":
        compact_pairs = []
        domain_counts = defaultdict(int)

        for pair in all_filtered_pairs:
            domain = pair[1]["section"]
            if domain_counts[domain] < MATRIX_TOPICS_PER_DOMAIN:
                compact_pairs.append(pair)
            else:
                hidden_count += 1
            domain_counts[domain] += 1

        if hidden_count:
            show_all = st.checkbox(
                (
                    f"显示全部研究主题（另有 {hidden_count} 项）"
                    if current_ui_language() == "zh-CN"
                    else f"Show all research topics ({hidden_count} more)"
                ),
                value=False,
                key="research_matrix_show_all",
            )
            if not show_all:
                display_pairs = compact_pairs

    filtered_rows = [pair[0] for pair in display_pairs]
    all_filtered_details = [
        pair[1]
        for pair in all_filtered_pairs
    ]

    st.caption(
        (
            "按 公司概览 → 商业化与竞争 → 财务与融资 → 团队与组织 → "
            "产品与技术 → 知识产权 自上而下阅读。"
            + (
                f" “全部”视图每个领域默认先显示 {MATRIX_TOPICS_PER_DOMAIN} 个主题；"
                "选择具体领域可查看完整目录。"
                if selected_section == "全部"
                else " 当前领域显示完整主题目录。"
            )
            if current_ui_language() == "zh-CN"
            else
            "Read top-down: Company Overview → Commercialization & Competition → "
            "Financials & Financing → Team & Organization → Product & Technology → "
            "Intellectual Property."
            + (
                f" The All view shows {MATRIX_TOPICS_PER_DOMAIN} topics per domain by default; select a domain to view its full topic list."
                if selected_section == "全部"
                else " The selected domain shows its full topic list."
            )
        )
    )
    st.dataframe(
        filtered_rows,
        use_container_width=True,
        hide_index=True,
        column_order=[
            "领域",
            "主题",
            "当前整理",
            "状态",
            "人工复核",
            "来源",
        ],
        column_config={
            "领域": ui("领域", "Domain"),
            "主题": ui("主题", "Topic"),
            "当前整理": ui("当前整理", "Current synthesis"),
            "状态": ui("状态", "AI status"),
            "人工复核": ui("人工复核", "Analyst review"),
            "来源": ui("来源", "Source"),
        },
    )

    selector_entries = []
    seen_labels = defaultdict(int)

    for item in all_filtered_details:
        base_label = (
            f"{domain_label(item['section'], current_ui_language())} · {item['label']}"
            if selected_section == "全部"
            else item["label"]
        )
        seen_labels[base_label] += 1
        selector_label = base_label
        if seen_labels[base_label] > 1:
            selector_label = (
                f"{base_label} ({seen_labels[base_label]})"
            )
        selector_entries.append(
            (selector_label, item)
        )

    labels = [
        ui("— 选择一个主题查看详情 —", "— Select a topic to view details —")
    ] + [
        label
        for label, _ in selector_entries
    ]

    if (
        st.session_state.get("research_matrix_topic")
        not in labels
    ):
        st.session_state["research_matrix_topic"] = labels[0]

    selected_label = st.selectbox(
        ui("主题详情", "Topic details"),
        labels,
        key="research_matrix_topic",
    )

    if selected_label.startswith("—"):
        return

    selected = next(
        item
        for label, item in selector_entries
        if label == selected_label
    )

    group = selected["group"]
    items = selected["items"]

    with st.container(border=True):
        st.markdown(
            f"**{domain_label(selected['section'], current_ui_language())} · "
            f"{clean_user_text(selected['label'])}**"
        )

        detail_text = clean_user_text(
            group.get("summary", "")
            or (items[0]["claim"] if items else "")
        )
        if detail_text:
            st.write(detail_text)

        st.caption(
            ui("AI 状态：", "AI status: ")
            + f"{research_status(items)}"
            + (
                ui(" · 来源：", " · Source: ") + compact_source_label(items)
                if items
                else ""
            )
            + (
                ui(" · ⚠ 证据锚点待检查", " · ⚠ Evidence anchor needs review")
                if evidence_anchor_warning(items)
                else ""
            )
        )

        render_analyst_review_controls(
            selected["review_key"]
        )

        with st.expander(ui("查看底层事实与原文证据", "View underlying facts and source evidence")):
            st.caption(
                ui(
                    "Evidence 只在需要核查时展开，不再单独占一个工作区页面。",
                    "Evidence expands only when needed for review; it no longer occupies a separate workspace view.",
                )
            )
            for item in items:
                st.markdown(
                    f"**{clean_user_text(item['claim'])}**"
                )
                st.caption(
                    f"{compact_location(item['source_location'])} · "
                    f"{item['evidence_status']} · "
                    + ui("信息归属：", "Attribution: ") + f"{item['attribution']}"
                )

                source_text = str(item.get("source_text", "")).strip()
                if source_text:
                    st.markdown(ui("原文证据：", "Source evidence:"))
                    st.write(source_text)

                check = item.get("evidence_check", {})
                if check.get("status") == "review":
                    st.warning(
                        ui("证据锚点检查：", "Evidence-anchor check: ")
                        + "；".join(check.get("issues", []))
                    )

def render_flags_tab(research_output, claim_map):
    entries = research_flag_entries(
        research_output,
        claim_map,
    )

    if not entries:
        st.success(ui("当前没有识别到需要优先人工核查的问题。", "No issue currently requires priority manual review."))
        return

    cleared_count = sum(
        1
        for entry in entries
        if entry["review"]["status"] == "cleared"
    )
    follow_up_count = sum(
        1
        for entry in entries
        if entry["review"]["status"] == "follow_up"
    )
    unreviewed_count = sum(
        1
        for entry in entries
        if entry["review"]["status"] == "unreviewed"
    )

    show_cleared = False
    if cleared_count:
        show_cleared = st.checkbox(
            (
                f"显示已人工处理的提醒（{cleared_count}）"
                if current_ui_language() == "zh-CN"
                else f"Show analyst-cleared flags ({cleared_count})"
            ),
            value=False,
            key="research_flags_show_cleared",
        )

    visible_entries = [
        entry
        for entry in entries
        if (
            entry["review"]["status"] != "cleared"
            or show_cleared
        )
    ]

    st.caption(
        (
            "Flags 只保留 Material + Unresolved + Actionable 的研究问题。"
            f" 未复核 {unreviewed_count} 项，人工标记需进一步跟进 {follow_up_count} 项，已处理 {cleared_count} 项。"
            "需要时直接在问题卡片里展开原文证据。"
            if current_ui_language() == "zh-CN"
            else
            "Flags keep only Material + Unresolved + Actionable research questions. "
            f"Unreviewed: {unreviewed_count}; analyst follow-up: {follow_up_count}; cleared: {cleared_count}. "
            "Expand source evidence directly inside each flag when needed."
        )
    )

    if not visible_entries:
        st.success(
            ui(
                "当前待处理的研究提醒已经清空。已处理项仍保留在本次会话的复核记录中。",
                "There are no open research flags. Cleared items remain in the session review record.",
            )
        )
        return

    for entry in visible_entries:
        flag = entry["flag"]
        related = entry["related"]
        review = get_analyst_review(
            entry["review_key"]
        )

        with st.container(border=True):
            title = clean_user_text(flag.get("title", ""))
            flag_type = clean_user_text(
                research_flag_type_label(
                    flag.get("type", "待核实")
                )
            )

            st.markdown(
                f"**{title}** · {flag_type}"
            )
            st.caption(
                ui("人工状态：", "Analyst status: ")
                + analyst_review_status_label(
                    review["status"]
                )
            )

            unresolved_question = clean_user_text(
                flag.get("unresolved_question", "")
            )
            if unresolved_question:
                st.markdown(
                    ui("**待回答：** ", "**Question:** ") + unresolved_question
                )

            why = clean_user_text(flag.get("why", ""))
            if why:
                st.write(why)

            action = clean_user_text(flag.get("action", ""))
            if action:
                st.caption(ui("建议：", "Next action: ") + action)

            if related:
                st.caption(
                    ui("涉及来源：", "Sources: ") + compact_source_label(related)
                )

                with st.expander(ui("查看相关事实与原文证据", "View related facts and source evidence")):
                    for item in related:
                        st.markdown(
                            f"**{clean_user_text(item['claim'])}**"
                        )
                        st.caption(
                            f"{compact_location(item['source_location'])} · "
                            f"{item['evidence_status']} · "
                            + ui("信息归属：", "Attribution: ") + f"{item['attribution']}"
                        )
                        source_text = str(
                            item.get("source_text", "")
                        ).strip()
                        if source_text:
                            st.write(source_text)

            render_analyst_review_controls(
                entry["review_key"]
            )

def render_sources_tab(claims):
    if not claims:
        st.caption("当前没有来源可展示。")
        return

    ordered = sorted(
        claims,
        key=source_sort_key,
    )

    grouped = defaultdict(list)
    for item in ordered:
        key = (
            item.get("source_filename", ""),
            item.get("source_location", ""),
        )
        grouped[key].append(item)

    filenames = {
        filename
        for filename, _ in grouped
        if filename
    }
    multi_file = len(filenames) > 1

    rows = []
    source_keys = []

    for (filename, location), items in grouped.items():
        categories = []
        for item in items:
            category = item.get("category", "其他")
            if category not in categories:
                categories.append(category)

        row = {
            "位置": compact_location(location),
            "Facts": len(items),
            "内容": " / ".join(categories[:2]),
            "状态": research_status(items),
        }

        if multi_file:
            stem = Path(filename).stem
            row = {
                "文件": stem[:18] + ("…" if len(stem) > 18 else ""),
                **row,
            }

        rows.append(row)
        source_keys.append((filename, location))

    st.caption(
        "Sources 是原始证据档案：Matrix / Flags 只保留事实与来源指针，"
        "需要核查原文、上下文或页面预览时再到这里展开。"
    )
    st.dataframe(
        rows,
        use_container_width=True,
        hide_index=True,
    )

    labels = [
        "— 选择来源查看详情 —"
    ] + [
        (
            f"{Path(filename).stem[:18]} · "
            if multi_file
            else ""
        )
        + compact_location(location)
        for filename, location in source_keys
    ]

    if (
        st.session_state.get("research_source_selector")
        not in labels
    ):
        st.session_state["research_source_selector"] = labels[0]

    selected_label = st.selectbox(
        "来源详情",
        labels,
        key="research_source_selector",
    )

    if selected_label.startswith("—"):
        return

    selected_index = labels.index(selected_label) - 1
    selected_key = source_keys[selected_index]
    selected_items = grouped[selected_key]

    with st.container(border=True):
        st.markdown(f"**{selected_label}**")
        st.caption(
            f"{len(selected_items)} 条底层事实 · "
            f"{research_status(selected_items)}"
        )

        for item in selected_items:
            st.markdown(
                f"**{clean_user_text(item['claim'])}**"
            )
            st.caption(
                f"{item['evidence_status']} · "
                + ui("信息归属：", "Attribution: ") + f"{item['attribution']}"
            )
            write_source_text_if_distinct(item)

        preview = next(
            (
                item.get("_preview_bytes")
                for item in selected_items
                if item.get("_preview_bytes")
            ),
            None,
        )

        if preview:
            with st.expander("查看原始页面预览"):
                st.image(
                    preview,
                    use_container_width=True,
                )

def render_research_output(research_output, claims):
    claim_map = {
        item["claim_id"]: item
        for item in claims
    }

    # Stateful navigation survives Streamlit reruns. Raw evidence is now
    # available in-place from Matrix / Flags instead of occupying a top-level tab.
    view_options = [
        "snapshot",
        "matrix",
        "flags",
    ]
    if (
        st.session_state.get("research_workspace_view")
        not in view_options
    ):
        st.session_state["research_workspace_view"] = view_options[0]

    selected_view = st.radio(
        ui("研究工作区", "Research workspace"),
        view_options,
        format_func=lambda value: {
            "snapshot": ui("Snapshot / 总览", "Snapshot"),
            "matrix": ui("Research Matrix / 研究矩阵", "Research Matrix"),
            "flags": ui("Flags / 研究提醒", "Flags"),
        }[value],
        horizontal=True,
        key="research_workspace_view",
        label_visibility="collapsed",
    )

    if selected_view == "snapshot":
        render_snapshot_tab(
            research_output,
            claims,
            claim_map,
        )
    elif selected_view == "matrix":
        render_matrix_tab(
            research_output,
            claim_map,
        )
    else:
        render_flags_tab(
            research_output,
            claim_map,
        )

def render_visual_evidence(item):
    visuals = item.get("visual_evidence", [])
    if not visuals:
        if item.get("image_count", 0) > 0:
            st.caption(
                f"该位置检测到 {item['image_count']} 个较大图片，但本次未形成视觉证据。"
            )
        return

    for index, visual in enumerate(visuals, start=1):
        summary = visual.get("summary", "")
        if summary:
            st.write(f"视觉 {index}：{summary}")
        for evidence in visual.get("evidence", []):
            if isinstance(evidence, dict) and evidence.get("text"):
                st.write(
                    f"- [{evidence.get('kind', 'other')}] {evidence['text']}"
                )


def render_review_flags(claims, review_flags):
    st.subheader("Needs Review / 待复核")

    if not review_flags:
        st.caption(
            "当前没有识别到需要重点复核的版本、口径或潜在冲突。"
        )
        return

    claim_map = {
        item["claim_id"]: item
        for item in claims
    }

    relation_labels = {
        "revision": "版本调整",
        "scope_difference": "口径 / 范围差异",
        "potential_conflict": "潜在冲突",
    }

    for flag in review_flags:
        with st.container(border=True):
            st.markdown(
                f"**{flag['comparison_key']}** · "
                f"{relation_labels.get(flag['relation'], flag['relation'])}"
            )
            st.caption(flag["reason"])

            for cid in flag["claim_ids"]:
                item = claim_map.get(cid)

                if item:
                    st.write(
                        f"• {item['claim']} — "
                        f"{item['source_filename']} · "
                        f"{item['source_location']}"
                    )


def render_claim_card(item):
    status_labels = {
        "Unverified": "未验证",
        "Supported": "当前材料支持",
        "Projection": "预测 / 目标",
        "Subjective / Marketing": "主观 / 营销表述",
    }

    kind_labels = {
        "text": "文字证据",
        "visual": "视觉证据",
        "mixed": "文字 + 视觉",
    }

    with st.container(border=True):
        left, right = st.columns([5, 1])

        with left:
            st.markdown(f"**{item['claim']}**")

        with right:
            st.caption(
                status_labels.get(
                    item["evidence_status"],
                    item["evidence_status"],
                )
            )

        st.caption(
            f"{item.get('category', '其他')} · "
            f"信息归属：{item['attribution']} · "
            f"{kind_labels.get(item.get('evidence_kind'), '—')}"
        )

        st.caption(
            f"来源：{item['source_filename']} · "
            f"{item['source_location']}"
        )

        with st.expander("查看证据"):
            st.markdown("**原始证据**")
            st.write(item.get("source_text", ""))

            context = item.get(
                "source_context",
                "",
            ).strip()

            if (
                context
                and context
                != item.get("source_text", "").strip()
            ):
                st.markdown("**附近上下文**")
                st.text(context)

            if item.get("visual_evidence"):
                st.markdown("**图片 / 图表读取结果**")
                render_visual_evidence(item)

            preview = item.get(
                "_preview_bytes"
            )

            if preview:
                st.markdown(
                    "**原始页面预览**"
                )
                st.image(
                    preview,
                    use_container_width=True,
                )


def render_claims_by_category(claims):
    st.subheader("Key Claims / 关键主张")

    by_category = defaultdict(list)

    for item in claims:
        by_category[
            item.get("category", "其他")
        ].append(item)

    category_order = [
        "客户与订单",
        "财务与预测",
        "商业化与产能",
        "技术与产品",
        "知识产权",
        "团队",
        "市场与竞争",
        "风险与限制",
        "公司概况",
        "其他",
    ]

    for category in category_order:
        items = by_category.get(category, [])

        if not items:
            continue

        st.markdown(
            f"### {category} ({len(items)})"
        )

        for item in items:
            render_claim_card(item)


def exportable_claims(claims):
    return [
        {
            key: value
            for key, value in item.items()
            if not key.startswith("_")
        }
        for item in claims
    ]


def diligence_source_refs(items):
    refs = []
    for item in items:
        source = (
            f"{item.get('source_filename', '')} · "
            f"{item.get('source_location', '')}"
        ).strip(" ·")
        if source and source not in refs:
            refs.append(source)
    return refs


def build_diligence_checklist_items(
    claims,
    research_output,
    analyst_reviews=None,
):
    analyst_reviews = analyst_reviews or {}
    claim_map = {
        item["claim_id"]: item
        for item in claims
    }

    checklist = []
    covered_fact_sets = []

    for flag in research_output.get(
        "research_flags",
        [],
    ):
        if not flag_is_actionable(
            flag,
            claim_map,
        ):
            continue

        related = [
            claim_map[fact_id]
            for fact_id in flag.get(
                "fact_ids",
                [],
            )
            if fact_id in claim_map
        ]
        review_key = flag_review_key(
            flag,
            related,
        )
        review = analyst_reviews.get(
            review_key,
            {
                "status": "unreviewed",
                "status_label": analyst_review_status_label(
                    "unreviewed"
                ),
                "note": "",
            },
        )

        if review.get("status") == "cleared":
            continue

        domain = canonical_matrix_domain(
            "",
            related,
        )
        fact_ids = tuple(sorted(
            item["claim_id"]
            for item in related
        ))
        covered_fact_sets.append(set(fact_ids))

        checklist.append({
            "kind": "Research Flag",
            "domain": domain,
            "title": clean_user_text(
                flag.get("title", "")
            ),
            "status": review.get(
                "status",
                "unreviewed",
            ),
            "status_label": review.get(
                "status_label",
                analyst_review_status_label(
                    review.get(
                        "status",
                        "unreviewed",
                    )
                ),
            ),
            "question": clean_user_text(
                flag.get(
                    "unresolved_question",
                    "",
                )
            ),
            "why": clean_user_text(
                flag.get("why", "")
            ),
            "action": clean_user_text(
                flag.get("action", "")
            ),
            "note": str(
                review.get("note", "")
                or ""
            ).strip(),
            "facts": [
                clean_user_text(
                    item.get("claim", "")
                )
                for item in related
            ],
            "sources": diligence_source_refs(
                related
            ),
            "fact_ids": list(fact_ids),
        })

    for section in research_output.get(
        "sections",
        [],
    ):
        for group in section.get(
            "groups",
            [],
        ):
            related = [
                claim_map[fact_id]
                for fact_id in group.get(
                    "fact_ids",
                    [],
                )
                if fact_id in claim_map
            ]
            if not related:
                continue

            domain = canonical_matrix_domain(
                section.get("title", ""),
                related,
            )
            label = canonical_matrix_topic(
                group.get("title", ""),
                related,
                domain,
            )
            review_key = matrix_review_key(
                domain,
                label,
                related,
            )
            review = analyst_reviews.get(
                review_key
            )

            if (
                not review
                or review.get("status")
                != "follow_up"
            ):
                continue

            fact_set = {
                item["claim_id"]
                for item in related
            }
            if any(
                fact_set == covered
                for covered in covered_fact_sets
            ):
                continue

            checklist.append({
                "kind": "Matrix Follow-up",
                "domain": domain,
                "title": clean_user_text(label),
                "status": "follow_up",
                "status_label": analyst_review_status_label(
                    "follow_up"
                ),
                "question": (
                    "该研究主题已被人工标记为需进一步跟进。"
                ),
                "why": clean_user_text(
                    group.get("summary", "")
                ),
                "action": (
                    "结合现有来源补充核验；"
                    "优先按人工备注执行。"
                ),
                "note": str(
                    review.get("note", "")
                    or ""
                ).strip(),
                "facts": [
                    clean_user_text(
                        item.get("claim", "")
                    )
                    for item in related
                ],
                "sources": diligence_source_refs(
                    related
                ),
                "fact_ids": sorted(fact_set),
            })

    checklist.sort(
        key=lambda item: (
            MATRIX_DOMAIN_ORDER.index(
                item["domain"]
            )
            if item["domain"]
            in MATRIX_DOMAIN_ORDER
            else len(MATRIX_DOMAIN_ORDER),
            0 if item["kind"] == "Research Flag" else 1,
            item["title"],
        )
    )

    return checklist


def make_diligence_checklist_pdf(
    claims,
    research_output,
    analyst_reviews=None,
):
    def safe(value):
        return html.escape(str(value or ""))

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pdfmetrics.registerFont(
        UnicodeCIDFont("STSong-Light")
    )
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Due Diligence Follow-up Checklist",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DDTitle",
        parent=styles["Title"],
        fontName="STSong-Light",
        fontSize=17,
        leading=22,
        spaceAfter=8,
    )
    h2 = ParagraphStyle(
        "DDH2",
        parent=styles["Heading2"],
        fontName="STSong-Light",
        fontSize=12,
        leading=17,
        spaceBefore=6,
        spaceAfter=5,
    )
    body = ParagraphStyle(
        "DDBody",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=9.5,
        leading=14,
        spaceAfter=4,
    )
    small = ParagraphStyle(
        "DDSmall",
        parent=body,
        fontSize=8.5,
        leading=12,
        textColor="#555555",
    )
    line = ParagraphStyle(
        "DDLine",
        parent=body,
        fontSize=9,
        leading=15,
        spaceAfter=2,
    )

    checklist = build_diligence_checklist_items(
        claims,
        research_output,
        analyst_reviews,
    )
    flag_count = sum(
        item["kind"] == "Research Flag"
        for item in checklist
    )
    manual_count = sum(
        item["kind"] == "Matrix Follow-up"
        for item in checklist
    )

    story = [
        Paragraph(
            "Due Diligence Follow-up Checklist / 尽调待核查清单",
            title_style,
        ),
        Paragraph(
            (
                f"待核查 {len(checklist)} 项 · "
                f"系统 Research Flags {flag_count} 项 · "
                f"人工标记 Matrix 跟进 {manual_count} 项"
            ),
            body,
        ),
        Paragraph(
            (
                "本清单只打印尚未处理的 Research Flags，"
                "以及研究员在 Matrix 中手动标记为“需进一步跟进”的主题。"
                "“已复核 / 暂无异议”的项目不会进入本清单。"
            ),
            small,
        ),
        Spacer(1, 5),
    ]

    if not checklist:
        story.append(
            Paragraph(
                "当前没有待核查事项。",
                body,
            )
        )

    for index, item in enumerate(
        checklist,
        start=1,
    ):
        story.append(
            Paragraph(
                (
                    f"{index}. [{safe(item['domain'])}] "
                    f"{safe(item['title'])}"
                ),
                h2,
            )
        )
        story.append(
            Paragraph(
                (
                    "[ ] 已核实　[ ] 需补充材料　"
                    "[ ] 继续访谈/外部核验　[ ] 暂不处理"
                ),
                body,
            )
        )
        story.append(
            Paragraph(
                (
                    "当前人工状态："
                    + safe(item["status_label"])
                    + (
                        "；备注：" + safe(item["note"])
                        if item.get("note")
                        else ""
                    )
                ),
                small,
            )
        )

        if item.get("question"):
            story.append(
                Paragraph(
                    "<b>待回答：</b>"
                    + safe(item["question"]),
                    body,
                )
            )
        if item.get("why"):
            story.append(
                Paragraph(
                    "<b>为什么重要：</b>"
                    + safe(item["why"]),
                    body,
                )
            )
        if item.get("action"):
            story.append(
                Paragraph(
                    "<b>建议核查：</b>"
                    + safe(item["action"]),
                    body,
                )
            )

        if item.get("facts"):
            story.append(
                Paragraph(
                    "<b>现有材料：</b>",
                    body,
                )
            )
            for fact in item["facts"][:5]:
                story.append(
                    Paragraph(
                        "• " + safe(fact),
                        small,
                    )
                )

        if item.get("sources"):
            story.append(
                Paragraph(
                    "<b>来源：</b>"
                    + safe(
                        "；".join(
                            item["sources"]
                        )
                    ),
                    small,
                )
            )

        story.extend([
            Paragraph(
                "核查结果：_______________________________________________",
                line,
            ),
            Paragraph(
                "________________________________________________________",
                line,
            ),
            Paragraph(
                "补充证据 / 联系人：________________________________________",
                line,
            ),
            Spacer(1, 7),
        ])

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def empty_run_stats():
    return {
        "api_calls": 0,
        "claim_calls": 0,
        "vision_calls": 0,
        "review_calls": 0,
        "triage_calls": 0,
        "synthesis_calls": 0,
        "verify_calls": 0,
    }


def build_live_run_key(
    provider,
    model_name,
    analyze_visuals,
    smoke_test,
    output_language,
    uploaded_files,
):
    if not uploaded_files:
        return None

    file_hashes = []
    for uploaded_file in uploaded_files:
        file_hashes.append(
            hashlib.sha256(
                uploaded_file.getvalue()
            ).hexdigest()
        )

    return hashlib.sha256(
        (
            PIPELINE_VERSION
            + "|"
            + provider
            + "|"
            + model_name
            + "|"
            + str(analyze_visuals)
            + "|"
            + str(smoke_test)
            + "|"
            + output_language
            + "|"
            + "|".join(file_hashes)
        ).encode("utf-8")
    ).hexdigest()


def render_live_workspace(result, current_run_key=None):
    claims = result["claims"]
    research_output = result["research_output"]
    analysis_counts = result["analysis_counts"]
    processing_warnings = list(
        result.get("warnings", [])
    )
    run_stats = result.get(
        "run_stats",
        empty_run_stats(),
    )
    cache_hit = bool(result.get("cache_hit"))
    smoke_test = bool(result.get("smoke_test"))
    elapsed_seconds = float(
        result.get("elapsed_seconds", 0)
    )
    stage_timings = dict(
        result.get("stage_timings", {})
    )

    if (
        current_run_key
        and result.get("run_key")
        and current_run_key != result["run_key"]
    ):
        st.info(
            ui(
                "当前显示的是上一次已完成的分析结果。你已经修改了文件或分析设置；点击“开始处理”后才会更新结果。",
                "The workspace is showing the previous completed result. Files or analysis settings have changed; click Start analysis to refresh it.",
            )
        )

    if cache_hit:
        st.info(
            ui(
                "已复用缓存结果，没有重新调用模型。",
                "Cached results were reused; the model was not called again.",
            )
        )

    if processing_warnings:
        with st.expander(
            (
                f"处理提示（{len(processing_warnings)}）"
                if current_ui_language() == "zh-CN"
                else f"Processing notes ({len(processing_warnings)})"
            ),
            expanded=False,
        ):
            for warning in processing_warnings:
                st.warning(warning)

    claim_map = {
        item["claim_id"]: item
        for item in claims
    }
    flag_entries = research_flag_entries(
        research_output,
        claim_map,
    )
    flags = [
        entry["flag"]
        for entry in flag_entries
        if entry["review"]["status"] != "cleared"
    ]

    runtime_text = (
        "<1s · cache"
        if cache_hit
        else f"{elapsed_seconds:.1f}s"
    )
    mode_text = ui(" · 快速测试", " · smoke test") if smoke_test else ""

    st.markdown(
        ui("**材料范围 ", "**Material scope ") + f"{analysis_counts['selected_units']}/"
        f"{analysis_counts['total_units']}**"
        f"{mode_text}"
        f"　·　Facts {len(claims)}"
        f"　·　Flags {len(flags)}"
        f"　·　{runtime_text}"
    )

    st.caption(
        ui("原始分析 API 调用：", "Model API calls: ")
        + f"{run_stats.get('api_calls', 0)} · "
        f"Triage {run_stats.get('triage_calls', 0)} · "
        f"Extract {run_stats.get('claim_calls', 0)} · "
        f"Vision {run_stats.get('vision_calls', 0)} · "
        f"Synthesis {run_stats.get('synthesis_calls', 0)} · "
        f"Verify {run_stats.get('verify_calls', 0)}"
    )

    if stage_timings:
        timing_parts = []
        for label, key in [
            ("Triage", "triage"),
            ("Extract", "extract"),
            ("Vision", "vision"),
            ("Synthesis", "synthesis"),
            ("Verify", "verify"),
        ]:
            seconds = float(
                stage_timings.get(key, 0) or 0
            )
            if seconds > 0:
                timing_parts.append(
                    f"{label} {seconds:.1f}s"
                )

        if timing_parts:
            st.caption(
                ui("阶段耗时：", "Stage timing: ")
                + " · ".join(timing_parts)
            )

    render_research_output(
        research_output,
        claims,
    )

    analyst_reviews = export_analyst_reviews()
    diligence_followups = build_diligence_checklist_items(
        claims,
        research_output,
        analyst_reviews,
    )
    export_data = {
        "research_output": research_output,
        "facts": exportable_claims(claims),
        "analysis_counts": analysis_counts,
        "analyst_reviews": analyst_reviews,
        "diligence_followups": diligence_followups,
    }

    json_bytes = json.dumps(
        export_data,
        ensure_ascii=False,
        indent=2,
    )

    try:
        pdf_bytes = make_diligence_checklist_pdf(
            claims,
            research_output,
            analyst_reviews,
        )
    except Exception as exc:
        pdf_bytes = None
        st.warning(f"尽调待核查清单 PDF 生成失败：{exc}")

    st.divider()
    st.subheader(ui("下载结果", "Downloads"))
    d1, d2 = st.columns(2)

    with d1:
        st.download_button(
            ui("下载 Research JSON", "Download Research JSON"),
            data=json_bytes,
            file_name="investment_research.json",
            mime="application/json",
            use_container_width=True,
        )

    with d2:
        if pdf_bytes:
            st.download_button(
                ui("下载尽调待核查清单 PDF", "Download diligence follow-up checklist PDF"),
                data=pdf_bytes,
                file_name="due_diligence_followup_checklist.pdf",
                mime="application/pdf",
                use_container_width=True,
            )


ui_language = st.selectbox(
    "Language / 语言",
    ["zh-CN", "en"],
    index=1,
    format_func=language_label,
    key="ui_language",
)

st.title("Investment Research AI")
st.caption("Evidence-first investment research workflow")
st.write(
    ui(
        "把长材料压缩成可快速阅读的研究总览、主题目录和少量研究提醒；重要结论保留来源，只有需要时才回到证据。",
        "Compress long investment materials into a fast research surface: a concise Snapshot, a structured Research Matrix and a small set of actionable Flags. Important conclusions stay traceable to source evidence.",
    )
)

mode = st.radio(
    ui("运行模式", "Mode"),
    ["Demo", "Live"],
    horizontal=True,
)

if mode == "Demo":
    st.info(
        ui(
            "Demo 使用虚构公司和合成尽调材料，无需 API Key；界面与 Live 共用同一套 Snapshot / Matrix / Flags；原文证据按需就地展开。",
            "Demo uses a fictional company and synthetic diligence materials. No API key is required; Demo and Live share the same Snapshot / Matrix / Flags workspace, with source evidence expanded only on demand.",
        )
    )

    if st.button(ui("运行 Demo", "Run Demo"), type="primary"):
        try:
            claims = build_claim_ledger(demo=True)
            normalize_demo_claims(claims)
            research_output = build_demo_research_output(
                claims
            )
            if ui_language == "en":
                claims, research_output = localize_demo_english(
                    claims,
                    research_output,
                )
        except Exception as exc:
            st.error(f"Demo 运行失败：{exc}")
            st.stop()

        st.session_state["active_demo_result"] = {
            "claims": claims,
            "research_output": research_output,
            "output_language": ui_language,
        }

    demo_result = st.session_state.get(
        "active_demo_result"
    )

    if demo_result:
        if demo_result.get("output_language") != ui_language:
            st.info(
                ui(
                    "语言已切换；重新运行 Demo 后内容会按当前语言重新生成。",
                    "Language changed; run Demo again to regenerate the synthetic result in the selected language.",
                )
            )
        claims = demo_result["claims"]
        research_output = demo_result["research_output"]
        claim_map = {
            item["claim_id"]: item
            for item in claims
        }
        visible_flags = [
            entry["flag"]
            for entry in research_flag_entries(
                research_output,
                claim_map,
            )
            if entry["review"]["status"] != "cleared"
        ]

        st.caption(
            f"Demo · Facts {len(claims)} · "
            f"Flags {len(visible_flags)} · "
            + ui("研究模块 ", "Research modules ")
            + f"{len(research_output.get('sections', []))}"
        )

        render_research_output(
            research_output,
            claims,
        )

        analyst_reviews = export_analyst_reviews()
        diligence_followups = build_diligence_checklist_items(
            claims,
            research_output,
            analyst_reviews,
        )
        demo_json = json.dumps(
            {
                "research_output": research_output,
                "facts": exportable_claims(claims),
                "analyst_reviews": analyst_reviews,
                "diligence_followups": diligence_followups,
            },
            ensure_ascii=False,
            indent=2,
        )

        try:
            demo_pdf = make_diligence_checklist_pdf(
                claims,
                research_output,
                analyst_reviews,
            )
        except Exception as exc:
            demo_pdf = None
            st.warning(f"Demo 尽调待核查清单 PDF 生成失败：{exc}")

        d1, d2 = st.columns(2)
        with d1:
            st.download_button(
                ui("下载 Demo Research JSON", "Download Demo Research JSON"),
                data=demo_json,
                file_name="demo_investment_research.json",
                mime="application/json",
                use_container_width=True,
            )
        with d2:
            if demo_pdf:
                st.download_button(
                    ui("下载 Demo 尽调待核查清单 PDF", "Download Demo diligence follow-up checklist PDF"),
                    data=demo_pdf,
                    file_name="demo_due_diligence_followup_checklist.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                )

else:
    with st.expander(
        ui("第一次使用？点这里看 30 秒教程", "First time here? Open the 30-second guide"),
        expanded=False,
    ):
        st.markdown(
            (
                """
**快速开始**

1. 选择模型提供商并输入对应 API Key。
2. 上传 TXT / PDF / DOCX / PPTX。
3. 普通短/中型材料（少于 41 个单元）直接 text-first；更长材料才先做模型筛选。
4. 改代码后建议先开“快速测试模式”，从材料前部、中部和尾部代表性抽样 5 个单元，避免一个格式问题烧完整份 API 费用。
5. “视觉补充”默认关闭；开启后最多补 2 个真正需要图表关系的页面。
6. 点“开始处理”。

完成后建议按这个顺序看：
- **Snapshot**：2 分钟看核心摘要、核心数字和需要注意的问题
- **Research Matrix**：看主题级当前整理、AI 状态、来源指针，并记录人工复核状态
- **Flags**：处理真正需要人工核查的问题；可标记“已复核 / 暂无异议”或“需进一步跟进”
- **Evidence**：不再单独占一个页面；需要时在 Matrix / Flags 中直接展开原文证据

系统的目标是减少需要阅读的内容，而不是把 PDF 拆成更多卡片。
API Key 只用于当前运行，不写入项目文件。
                """
                if ui_language == "zh-CN"
                else
                """
**Quick start**

1. Select a model provider and enter your own API key.
2. Upload TXT / PDF / DOCX / PPTX materials.
3. Short and medium materials (fewer than 41 units) go directly to text-first extraction; longer materials use model triage first.
4. After code changes, use Smoke Test first: five representative units are sampled across the front, middle and end of selected material.
5. Visual follow-up is off by default; when enabled, at most two relationship-heavy or text-insufficient pages are inspected.
6. Click **Start analysis**.

Recommended reading order:
- **Snapshot** — the two-minute view of key points, core numbers and issues that matter
- **Research Matrix** — topic-level synthesis, AI status, source pointers and analyst review state
- **Flags** — only material, unresolved and actionable diligence questions
- **Evidence** — no separate top-level page; expand original evidence inside Matrix / Flags when needed

The goal is to reduce reading, not to split a PDF into more cards.
The API key is used only for the current run and is never written into project files.
                """
            )
        )

    col1, col2 = st.columns(2)
    with col1:
        provider = st.selectbox(
            ui("模型提供商", "Model provider"),
            list(PROVIDERS.keys()),
            index=0,
        )
    with col2:
        model_name = st.text_input(
            ui("模型", "Model"),
            value=PROVIDERS[provider]["default_model"],
            help=ui("默认模型已经填好；一般不需要修改。", "A default model is prefilled; you normally do not need to change it."),
        )

    api_key = st.text_input(
        PROVIDERS[provider]["key_label"],
        type="password",
        help=ui("密钥只用于当前运行，不写入文件或结果。", "The key is used only for the current run and is not written to files or results."),
    )

    analyze_visuals = st.checkbox(
        ui("必要时补充读取重要图表 / 截图（最多 2 个材料单元）", "Use visual follow-up when needed (max 2 material units)"),
        value=False,
        help=ui(
            "默认关闭以节约时间和费用。开启后也不会逐页 Vision；只补文字明显不足，或竞争格局/客户/股权/融资时间线等关系图确实影响判断的页面。",
            "Off by default to save time and cost. Even when enabled, Vision is not run page-by-page; it is limited to text-insufficient or relationship-heavy pages that materially affect interpretation.",
        ),
    )

    smoke_test = st.checkbox(
        ui("快速测试模式（每份文件代表性抽样 5 个材料单元）", "Smoke Test (5 representative material units per file)"),
        value=False,
        help=ui(
            "用于改代码后的低成本验证。会从筛选后的材料前部、中部和尾部均匀抽样 5 个单元，避免只测试前几页。",
            "Low-cost validation after code changes. Samples five units across the front, middle and end of selected material rather than testing only the first pages.",
        ),
    )

    uploaded_files = st.file_uploader(
        ui("上传研究材料", "Upload research materials"),
        type=SUPPORTED_TYPES,
        accept_multiple_files=True,
        help=ui(
            "支持 TXT、PDF、DOCX 和 PPTX。少于 41 个材料单元时直接 text-first 分析；更长材料才启用模型筛选。",
            "Supports TXT, PDF, DOCX and PPTX. Materials with fewer than 41 units go directly to text-first analysis; longer materials use model triage.",
        ),
    )

    if st.button(ui("开始处理", "Start analysis"), type="primary"):
        if not api_key:
            st.error(ui("请输入 API Key。", "Please enter an API key."))
            st.stop()
        if not uploaded_files:
            st.error(ui("请至少上传一份材料。", "Please upload at least one file."))
            st.stop()

        try:
            client = create_client(provider, api_key)
        except Exception as exc:
            st.error(f"创建模型客户端失败：{exc}")
            st.stop()

        reset_run_stats()
        reset_stage_timings()
        started_at = time.perf_counter()

        file_payloads = [
            (uploaded_file, uploaded_file.getvalue())
            for uploaded_file in uploaded_files
        ]

        run_cache_key = build_live_run_key(
            provider,
            model_name,
            analyze_visuals,
            smoke_test,
            ui_language,
            uploaded_files,
        )

        result_cache = st.session_state.setdefault(
            "result_cache",
            {},
        )

        if run_cache_key in result_cache:
            cached = result_cache[run_cache_key]
            claims = cached["claims"]
            research_output = cached["research_output"]
            analysis_counts = cached["analysis_counts"]
            processing_warnings = list(cached.get("warnings", []))
            cached_run_stats = cached.get(
                "run_stats",
                empty_run_stats(),
            )
            cached_elapsed_seconds = float(
                cached.get("elapsed_seconds", 0)
            )
            cached_stage_timings = cached.get(
                "stage_timings",
                {},
            )
            cache_hit = True
        else:
            cache_hit = False
            claims = []
            processing_warnings = []
            rate_limit_hit = False
            analysis_counts = {
                "total_units": 0,
                "selected_units": 0,
                "skipped_units": 0,
                "vision_units": 0,
            }
            progress = st.progress(0)
            status = st.empty()

            try:
                for file_index, (uploaded_file, file_bytes) in enumerate(
                    file_payloads,
                    start=1,
                ):
                    status.write(f"快速解析：{uploaded_file.name}")

                    units = extract_document_units(
                        file_bytes,
                        uploaded_file.name,
                        analyze_visuals=analyze_visuals,
                        client=client,
                        model_name=model_name,
                        status_callback=status.write,
                    )

                    analysis_counts["total_units"] += len(units)

                    try:
                        status.write(
                            f"筛选真正值得分析的内容：{uploaded_file.name}"
                        )
                        stage_started = time.perf_counter()
                        selected_units, skipped_units = triage_units(
                            client,
                            units,
                            model_name,
                        )
                        add_stage_time(
                            "triage",
                            time.perf_counter() - stage_started,
                        )
                    except Exception as exc:
                        if is_rate_limit_error(exc):
                            rate_limit_hit = True
                            processing_warnings.append(
                                "页面筛选阶段 API 达到限额。"
                            )
                            break
                        selected_units = units
                        skipped_units = []
                        processing_warnings.append(
                            f"{uploaded_file.name}：页面筛选失败，"
                            f"已保守分析全部内容。原因：{exc}"
                        )

                    if smoke_test and len(selected_units) > 5:
                        sampled_units, omitted_units = (
                            representative_sample_units(
                                selected_units,
                                limit=5,
                            )
                        )
                        selected_units = sampled_units
                        skipped_units = (
                            skipped_units
                            + omitted_units
                        )
                        sampled_locations = " / ".join(
                            unit["source_location"]
                            for unit in selected_units
                        )
                        processing_warnings.append(
                            f"{uploaded_file.name}：快速测试模式代表性抽样 5 个材料单元："
                            f"{sampled_locations}"
                        )

                    analysis_counts["selected_units"] += len(selected_units)
                    analysis_counts["skipped_units"] += len(skipped_units)

                    if analyze_visuals:
                        try:
                            stage_started = time.perf_counter()
                            visual_count = enrich_selected_visuals(
                                client,
                                selected_units,
                                model_name,
                                enabled=True,
                                status_callback=status.write,
                            )
                            add_stage_time(
                                "vision",
                                time.perf_counter() - stage_started,
                            )
                            analysis_counts["vision_units"] += visual_count
                        except Exception as exc:
                            processing_warnings.append(
                                f"{uploaded_file.name}：视觉补充失败，"
                                f"继续使用文字结果。原因：{exc}"
                            )

                    source_type = (
                        Path(uploaded_file.name)
                        .suffix
                        .lower()
                        .lstrip(".")
                    )

                    batches = [
                        selected_units[i:i + BATCH_SIZE]
                        for i in range(
                            0,
                            len(selected_units),
                            BATCH_SIZE,
                        )
                    ]

                    for batch_index, batch in enumerate(
                        batches,
                        start=1,
                    ):
                        locations = (
                            f"{batch[0]['source_location']} - "
                            f"{batch[-1]['source_location']}"
                        )
                        status.write(
                            f"提取高价值事实：{uploaded_file.name} · "
                            f"{locations} ({batch_index}/{len(batches)})"
                        )

                        try:
                            stage_started = time.perf_counter()
                            batch_claims, batch_warnings = (
                                extract_claims_from_batch(
                                    client,
                                    batch,
                                    uploaded_file.name,
                                    source_type,
                                    model_name,
                                )
                            )
                            add_stage_time(
                                "extract",
                                time.perf_counter() - stage_started,
                            )
                            claims.extend(batch_claims)

                            for warning in batch_warnings:
                                processing_warnings.append(
                                    f"{uploaded_file.name} · "
                                    f"{locations}：{warning}"
                                )

                        except Exception as exc:
                            if is_rate_limit_error(exc):
                                processing_warnings.append(
                                    f"{uploaded_file.name} · {locations}："
                                    "API 已达到限额，已停止后续模型调用。"
                                )
                                rate_limit_hit = True
                                break

                            processing_warnings.append(
                                f"{uploaded_file.name} · {locations}："
                                f"{exc}，已跳过该批次。"
                            )

                    if rate_limit_hit:
                        break

                    progress.progress(
                        file_index / len(file_payloads)
                    )

            except Exception as exc:
                st.error(f"文件处理失败：{exc}")
                st.stop()

            add_claim_ids(claims)

            if not claims:
                status.empty()
                progress.empty()

                if processing_warnings:
                    with st.expander(
                        f"查看处理警告（{len(processing_warnings)}）",
                        expanded=True,
                    ):
                        for warning in processing_warnings:
                            st.warning(warning)

                st.warning("没有提取到可用于研究汇总的高价值事实。")
                st.stop()

            if rate_limit_hit:
                research_output = fallback_research_output(claims, output_language=ui_language)
                processing_warnings.append(
                    "API 额度已用完；本次使用已完成事实生成本地降级版汇总。"
                )
            else:
                status.write("生成研究总览、目录和主题聚合…")
                try:
                    stage_started = time.perf_counter()
                    research_output = synthesize_research_output(
                        client,
                        claims,
                        model_name,
                        partial_scope=smoke_test,
                        output_language=ui_language,
                    )
                    research_output, suppressed_flags = (
                        apply_partial_scope_flag_guard(
                            research_output,
                            smoke_test,
                        )
                    )
                    if suppressed_flags:
                        processing_warnings.append(
                            "快速测试仅覆盖部分材料：已抑制 "
                            f"{suppressed_flags} 个依赖未抽样页面的缺失型 Flag。"
                        )

                    research_output, added_caveat_flags = (
                        ensure_material_caveat_flags(
                            research_output,
                            claims,
                            partial_scope=smoke_test,
                            output_language=ui_language,
                        )
                    )
                    if added_caveat_flags:
                        processing_warnings.append(
                            "一致性检查补入 "
                            f"{added_caveat_flags} 个已在摘要/Matrix中明确暴露的未解决问题。"
                        )

                    add_stage_time(
                        "synthesis",
                        time.perf_counter() - stage_started,
                    )
                except Exception as exc:
                    processing_warnings.append(
                        f"研究汇总生成失败，已使用本地降级版：{exc}"
                    )
                    research_output = fallback_research_output(claims, output_language=ui_language)

                status.write("验证最终总结是否被证据支持…")
                try:
                    stage_started = time.perf_counter()
                    research_output = verify_research_output(
                        client,
                        research_output,
                        claims,
                        model_name,
                        output_language=ui_language,
                    )
                    add_stage_time(
                        "verify",
                        time.perf_counter() - stage_started,
                    )
                except Exception as exc:
                    processing_warnings.append(
                        f"最终语义验证失败，保留第一版汇总：{exc}"
                    )

                if (
                    not research_output.get("executive_summary")
                    and not research_output.get("sections")
                ):
                    processing_warnings.append(
                        "语义验证后没有保留足够内容，已回退到证据约束的本地聚合结果。"
                    )
                    research_output = fallback_research_output(claims, output_language=ui_language)

            research_output = merge_computed_metrics(
                research_output,
                claims,
                output_language=ui_language,
            )

            status.empty()
            progress.empty()

        elapsed_seconds = time.perf_counter() - started_at

        if cache_hit:
            display_run_stats = cached_run_stats
            display_elapsed_seconds = cached_elapsed_seconds
            display_stage_timings = dict(
                cached_stage_timings
            )
        else:
            display_run_stats = dict(get_run_stats())
            display_elapsed_seconds = elapsed_seconds
            display_stage_timings = get_stage_timings()

        result_payload = {
            "run_key": run_cache_key,
            "claims": claims,
            "research_output": research_output,
            "analysis_counts": analysis_counts,
            "warnings": processing_warnings,
            "run_stats": display_run_stats,
            "elapsed_seconds": display_elapsed_seconds,
            "stage_timings": display_stage_timings,
            "cache_hit": cache_hit,
            "smoke_test": smoke_test,
            "output_language": ui_language,
        }

        st.session_state["active_live_result"] = result_payload

        if not cache_hit:
            result_cache[run_cache_key] = {
                "claims": claims,
                "research_output": research_output,
                "analysis_counts": analysis_counts,
                "warnings": processing_warnings,
                "run_stats": display_run_stats,
                "elapsed_seconds": display_elapsed_seconds,
                "stage_timings": display_stage_timings,
                "output_language": ui_language,
            }

    current_run_key = build_live_run_key(
        provider,
        model_name,
        analyze_visuals,
        smoke_test,
        ui_language,
        uploaded_files,
    )

    active_live_result = st.session_state.get(
        "active_live_result"
    )

    if active_live_result:
        render_live_workspace(
            active_live_result,
            current_run_key=current_run_key,
        )

st.divider()
st.caption(
    ui(
        "系统优先减少阅读量，并确保最终总结可回到证据。它验证的是“是否忠实于材料”，不是材料本身是否真实；事实真伪仍需要合同、访谈或第三方来源等进一步尽调。",
        "The system prioritizes reading reduction while keeping final synthesis traceable to evidence. It checks faithfulness to supplied materials, not whether the materials are objectively true; factual verification still requires contracts, interviews or third-party diligence.",
    )
)
